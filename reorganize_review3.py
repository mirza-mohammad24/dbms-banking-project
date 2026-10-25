"""Reorganize the manually edited Review 3 using lossless OOXML block moves.

Never rebuild from the original template; Review 2 is read-only throughout.
"""
from copy import deepcopy
from collections import Counter
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile
import hashlib
import json
import re
import sys
import posixpath

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from lxml import etree

ROOT = Path(__file__).resolve().parent
TARGET = ROOT / 'Project Template - review 3.docx'
GUIDE = ROOT / 'Guide.docx'
OTHER = ROOT / 'Project Template - review 2.docx'
WORK = ROOT / '.report_work' / 'review3-reorganization'
WORK.mkdir(parents=True, exist_ok=True)


def text(e):
    return ''.join(e.xpath('.//w:t/text()'))


def style(e):
    s = e.xpath('./w:pPr/w:pStyle/@w:val')
    return s[0] if s else ''


def set_text(e, value):
    p = Paragraph(e, None)
    p.clear()
    p.add_run(value)


def para(value, sty='Normal', page=False):
    e = OxmlElement('w:p')
    p = Paragraph(e, None)
    pp = e.get_or_add_pPr()
    s = OxmlElement('w:pStyle')
    s.set(qn('w:val'), sty)
    pp.append(s)
    if page:
        pb = OxmlElement('w:pageBreakBefore')
        pp.append(pb)
    p.add_run(value)
    return e


def rename(e, value, sty=None, page=False):
    set_text(e, value)
    pp = e.get_or_add_pPr()
    if sty:
        ps = pp.find(qn('w:pStyle'))
        if ps is None:
            ps = OxmlElement('w:pStyle')
            pp.insert(0, ps)
        ps.set(qn('w:val'), sty)
    for old in list(pp.findall(qn('w:pageBreakBefore'))):
        pp.remove(old)
    if page:
        pp.append(OxmlElement('w:pageBreakBefore'))


def write_package(source, output, document):
    xml = etree.tostring(document.element, xml_declaration=True,
                         encoding='UTF-8', standalone=True)
    tmp = output.with_suffix('.reorganizing.docx')
    with ZipFile(source) as src, ZipFile(tmp, 'w') as dst:
        for info in src.infolist():
            dst.writestr(info, xml if info.filename == 'word/document.xml'
                         else src.read(info.filename))
    tmp.replace(output)


def fingerprint(nodes):
    return [etree.tostring(e, method='c14n', exclusive=True) for e in nodes]


def images(nodes):
    return Counter(r for e in nodes for r in e.xpath('.//a:blip/@r:embed'))


def refresh_toc(pdf):
    from pypdf import PdfReader
    document = Document(TARGET)
    pages = [p.extract_text() for p in PdfReader(pdf).pages]
    norm = lambda t: re.sub(r'\s+', ' ', t).strip().casefold()
    mapping = {}
    for p in document.paragraphs:
        if style(p._p) != 'ReportTOC':
            continue
        heading = p.text.rsplit('\t', 1)[0]
        found = [i + 1 for i, page in enumerate(pages)
                 if any(norm(line) == norm(heading) for line in page.splitlines())]
        if not found:
            raise RuntimeError('Missing rendered heading: ' + heading)
        mapping[heading] = found[-1]
        set_text(p._p, heading + '\t' + str(found[-1]))
    write_package(TARGET, TARGET, document)
    (WORK / 'page_map.json').write_text(json.dumps(mapping, indent=2), encoding='utf-8')
    print(json.dumps({'pages': len(pages), 'page_map': mapping}, indent=2))


def polish_guide():
    """Keep questions with their answers and omit unused copied screenshot data."""
    with ZipFile(GUIDE) as src:
        parts = {n: src.read(n) for n in src.namelist()}
        infos = src.infolist()
    document = Document(GUIDE)
    in_viva = False
    for p in document.paragraphs:
        if p.text == '4 Viva Questions and Answers':
            in_viva = True
        if in_viva:
            p.paragraph_format.keep_together = True
        if re.match(r'^\d+\. ', p.text):
            p.paragraph_format.keep_with_next = True
    parts['word/document.xml'] = etree.tostring(document.element, xml_declaration=True,
                                               encoding='UTF-8', standalone=True)
    rns = '{http://schemas.openxmlformats.org/package/2006/relationships}'
    rels = etree.fromstring(parts['word/_rels/document.xml.rels'])
    used = set(document.element.xpath('//@r:embed | //@r:id | //@r:link'))
    for rel in list(rels):
        if rel.get('Type', '').endswith('/image') and rel.get('Id') not in used:
            rels.remove(rel)
    parts['word/_rels/document.xml.rels'] = etree.tostring(rels, xml_declaration=True,
                                                         encoding='UTF-8', standalone=True)
    referenced = set()
    for name, raw in parts.items():
        if name.endswith('.rels'):
            base = posixpath.dirname(posixpath.dirname(name))
            for rel in etree.fromstring(raw):
                if rel.get('TargetMode') != 'External':
                    referenced.add(posixpath.normpath(posixpath.join(base, rel.get('Target', ''))))
    orphan_media = {n for n in parts if n.startswith('word/media/') and n not in referenced}
    for name in list(parts):
        if name.startswith('word/header') and name.endswith('.xml'):
            header = etree.fromstring(parts[name])
            ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
            texts = header.xpath('//w:t', namespaces=ns)
            if texts:
                texts[0].text = 'BCSE302P CORE BANKING PROJECT GUIDE'
                for t in texts[1:]:
                    t.text = ''
            parts[name] = etree.tostring(header, xml_declaration=True, encoding='UTF-8', standalone=True)
    content_types = etree.fromstring(parts['[Content_Types].xml'])
    for e in list(content_types):
        if e.get('PartName', '').lstrip('/') in orphan_media:
            content_types.remove(e)
    parts['[Content_Types].xml'] = etree.tostring(content_types, xml_declaration=True,
                                                encoding='UTF-8', standalone=True)
    temp = GUIDE.with_suffix('.polishing.docx')
    with ZipFile(temp, 'w') as dst:
        for info in infos:
            if info.filename not in orphan_media:
                dst.writestr(info, parts[info.filename])
    temp.replace(GUIDE)
    print(json.dumps({'unused_guide_media_removed': len(orphan_media),
                      'guide_bytes': GUIDE.stat().st_size}))


def reorganize():
    if GUIDE.exists():
        raise RuntimeError('Guide.docx already exists; inspect before overwriting.')
    other_hash = hashlib.sha256(OTHER.read_bytes()).hexdigest()
    backup = ROOT / 'report_backups' / 'before_review3_reorganization'
    backup.mkdir(parents=True, exist_ok=True)
    source = backup / (datetime.now().strftime('%Y%m%d_%H%M%S_') + TARGET.name)
    source.write_bytes(TARGET.read_bytes())
    document = Document(source)
    original = list(document.element.body)

    def at(value):
        return next(i for i, e in enumerate(original)
                    if style(e).startswith('Heading') and text(e) == value)

    checklist_start = at('8.9 Review 2 Requirements Checklist')
    work_start = at('9 Work break down')
    requirements_start = at('10 Review Requirements')
    concurrency_start = at('10.1 Concurrency and Locking Demonstration')
    tests_start = at('10.2 Comprehensive Verification Test Suite')
    runbook_start = at('10.3 Live Demonstration Runbook')
    app_start = at('10.4 Working Application Modules')
    viva_start = at('10.5 Viva Preparation and Evaluation Anchors')
    sect = next(e for e in original if e.tag == qn('w:sectPr'))
    end = original.index(sect)
    checklist = deepcopy(original[checklist_start:work_start])
    work = deepcopy(original[work_start:requirements_start])
    concurrency = deepcopy(original[concurrency_start:tests_start])
    tests = deepcopy(original[tests_start:runbook_start])
    runbook = deepcopy(original[runbook_start:app_start])
    app = deepcopy(original[app_start:viva_start])
    viva = deepcopy(original[viva_start:end])
    prefix = deepcopy(original[:checklist_start])

    rename(concurrency[0], '8.9 Concurrency and Locking Demonstration')
    rename(tests[0], '8.10 Testing and Validation')
    rename(app[0], '9 Application Screenshots', 'Heading1', page=True)
    rename(work[0], '10 Work Breakdown', 'Heading1', page=True)

    # Only this overview sentence changes outside the moved sections and TOC.
    for e in prefix:
        if 'a timed demonstration runbook and complete application' in text(e):
            p = Paragraph(e, None)
            for run in p.runs:
                run.text = run.text.replace('a timed demonstration runbook and complete application',
                                            'complete application')
            # Runs can split the sentence in manually edited Word documents.
            if 'a timed demonstration runbook' in text(e):
                set_text(e, text(e).replace('a timed demonstration runbook and complete application',
                                           'complete application'))

    body = document.element.body
    for e in list(body):
        body.remove(e)
    for e in prefix + concurrency + tests + app + work + [deepcopy(sect)]:
        body.append(e)

    # Rebuild the static contents entries to match the final section structure.
    old_toc = [e for e in body if style(e) == 'ReportTOC']
    template = deepcopy(old_toc[0])
    toc_index = list(body).index(old_toc[0])
    for e in old_toc:
        body.remove(e)
    headings = [text(e) for e in body if
                (style(e) == 'Heading1' and re.match(r'^\d+ ', text(e))) or
                (style(e) == 'Heading2' and re.match(r'^8\.\d+ ', text(e)))]
    for offset, heading in enumerate(headings):
        e = deepcopy(template)
        set_text(e, heading + '\t0')
        body.insert(toc_index + offset, e)

    # A package clone preserves all styles and relationships of moved content.
    guide = Document(source)
    gb = guide.element.body
    for e in list(gb):
        gb.remove(e)
    rename(checklist[0], '1 Database Completion Checklist', 'Heading1')
    rename(runbook[0], '3 Live Demonstration Runbook', 'Heading1', page=True)
    rename(viva[0], '4 Viva Questions and Answers', 'Heading1', page=True)
    guide_nodes = [
        para('Core Banking Project Guide', 'Title'),
        para('Review Preparation and Live Demonstration', 'Subtitle'),
        para('This guide brings together the database completion checklist, faculty review '
             'requirements, live demonstration steps and project-specific viva answers. '
             'Use it to prepare the demonstration and individual viva. The project report '
             'contains the implementation, test evidence and application screenshots.'),
    ] + checklist + [
        para('2 Faculty Review Reference', 'Heading1', page=True),
        para('Required deliverables', 'Heading2'),
        para('Complete populated database; full scripts that can be rerun from scratch; '
             'advanced SQL; views; procedures; functions; triggers; transactions and '
             'concurrency; indexing and query optimization; security implementation; '
             'at least fifteen test cases; final report; a ten to twelve minute live '
             'demonstration and individual viva.'),
        para('The faculty also asks for the design and screenshots of all working '
             'application modules, showing accessible inputs and outputs, and indexing '
             'to improve a search query.'),
        para('Required demonstration flow', 'Heading2'),
        para('1. Run DDL and DML on a clean database.'),
        para('2. Demonstrate role based access.'),
        para('3. Show a transactional workflow with a deliberate failure and rollback.'),
        para('4. Show a trigger firing.'),
        para('5. Compare an optimization before and after indexing.'),
        para('6. Demonstrate one concurrency scenario.'),
        para('Evaluation guidance', 'Heading2'),
        para('Excellent at 85 percent or above: justified design, verified SQL, '
             'reasoned optimization, demonstrated anomalies and fluent viva. '
             'Good at 70 to 84 percent: correct and complete work with limited '
             'justification or shallow optimization. Satisfactory at 50 to 69 percent: '
             'working implementation with thin design or analysis. Poor below 50 '
             'percent: missing components, errors or unexplained work.'),
        para('The evaluator chooses five or six questions per student. Each answer '
             'is marked from zero to three: zero for an incorrect answer, one for '
             'a definition alone, two for a correct answer with an example, and '
             'three for a correct, justified answer connected to this project.'),
    ] + runbook + viva + [deepcopy(sect)]
    for e in guide_nodes:
        gb.append(e)

    # Assert that every original block is retained in its destination unchanged,
    # apart from the explicitly renamed headings and the short overview update.
    assert fingerprint(concurrency[1:]) == fingerprint(original[concurrency_start+1:tests_start])
    assert fingerprint(tests[1:]) == fingerprint(original[tests_start+1:runbook_start])
    assert fingerprint(app[1:]) == fingerprint(original[app_start+1:viva_start])
    assert fingerprint(work[1:]) == fingerprint(original[work_start+1:requirements_start])
    assert fingerprint(runbook[1:]) == fingerprint(original[runbook_start+1:app_start])
    assert fingerprint(viva[1:]) == fingerprint(original[viva_start+1:end])
    s6 = at('6 Tables and Constraints')
    s7 = at('7 SQL Implementation')
    s8 = at('8 Advanced SQL Implementation')
    assert fingerprint(prefix[s6:s7]) == fingerprint(original[s6:s7])
    assert fingerprint(prefix[s8:]) == fingerprint(original[s8:checklist_start])
    assert images(original) == images(list(body)) + images(checklist + runbook + viva)
    assert not any('Review Requirements' == text(e).split(' ', 1)[-1] for e in body)
    assert sum(bool(re.match(r'^\d+\. ', text(e))) for e in viva[1:]) == 31

    write_package(source, TARGET, document)
    write_package(source, GUIDE, guide)
    assert other_hash == hashlib.sha256(OTHER.read_bytes()).hexdigest()
    with ZipFile(source) as a, ZipFile(TARGET) as b:
        assert all(a.read(n) == b.read(n) for n in a.namelist() if n != 'word/document.xml')
    audit = {
        'backup': str(source), 'review2_sha256': other_hash,
        'main_images_before': sum(images(original).values()),
        'main_images_after': sum(images(list(body)).values()),
        'guide_moved_images': sum(images(checklist + runbook + viva).values()),
        'viva_answers_moved': 31,
        'section6_exact_xml_preserved': True,
        'existing_section8_exact_xml_preserved': True,
        'all_non_document_package_parts_preserved': True,
        'main_headings': headings,
    }
    (WORK / 'audit.json').write_text(json.dumps(audit, indent=2), encoding='utf-8')
    print(json.dumps(audit, indent=2))


if __name__ == '__main__':
    if '--polish-guide' in sys.argv:
        polish_guide()
    elif '--refresh-toc' in sys.argv:
        refresh_toc(Path(sys.argv[sys.argv.index('--refresh-toc')+1]))
    else:
        reorganize()
