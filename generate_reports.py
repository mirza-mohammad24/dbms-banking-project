"""Populate both VIT review reports from retained templates and captured SQL.

Run: uv run python generate_reports.py [--refresh-evidence] [--page-map FILE]
Original templates are retained in report_backups/original. Live outputs are
collected by collect_report_evidence.py; it never resets banking_db. Rebuilding
is deterministic for a given evidence.json and always starts from the originals.
"""
from __future__ import annotations

import argparse
from copy import deepcopy
from datetime import datetime
import hashlib
import json
from pathlib import Path
import re
import shutil
import subprocess
from zipfile import ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from collect_report_evidence import split_sql, clean_sql, TABLES

ROOT=Path(__file__).resolve().parent
WORK=ROOT/'.report_work'
TITLE='CORE BANKING & TRANSACTION MANAGEMENT SYSTEM'
MEMBERS=[('MIRZA MOHAMMAD ABBAS','24BDS0202','9451792098'),
         ('ADITYA KUMAR SAHU','24BAI0128','7870457312'),
         ('RISHAV MANJHI','24BAI0001','6287568398')]

TABLE_INFO={
 'Postal_Directory':('Canonical postal location lookup shared by branches.',
   'Pincode determines City and State. Both are atomic facts about the whole postal key; neither City nor State determines a unique pincode in this model.'),
 'Bank_Directory':('Bank-name lookup for a beneficiary destination IFSC.',
   'IFSC_Code determines BankName. A bank name may have many IFSC codes, so BankName is not a determinant of any other non-key fact in this relation.'),
 'Branch':('Servicing branch identity, contact details and postal reference.',
   'BranchID determines the branch facts; IFSC_Code and ContactPhone are additional unique candidate keys. City and State are looked up through Pincode instead of being repeated here.'),
 'Customer':('KYC identity and contact master for a bank customer.',
   'CustomerID determines the identity facts. Phone, Email and PAN_Number are declared unique candidate keys. No non-key determinant of another independent customer fact is assumed.'),
 'Account':('A single-owner Savings, Current or Salary account at one branch.',
   'AccountNo determines CustomerID, BranchID and account-specific facts. Owner and branch descriptions remain in their own relations, removing identity and location transitive dependencies.'),
 'Transaction_History':('Positive-valued account events with channel, narrative and timestamp.',
   'TransactionID determines one event. AccountNo can repeat and does not determine Amount, PaymentMode or TransactionDate; account and owner facts are not copied into the event.'),
 'Loan':('Borrower and issuing-branch loan contract, debt state and generated EMI.',
   'LoanID determines the independently stored contract and state facts. Borrower and branch descriptions are separated. EMI_Amount is an enforced generated projection of PrincipalAmount, InterestRate and TenureMonths; see the formal derived-column qualification below.'),
 'Loan_Repayment':('A loan payment attempt, including pending and failed attempts.',
   'PaymentID determines LoanID and attempt-specific facts. LoanID can repeat across installments; loan terms and borrower identity are not redundantly stored with each attempt.'),
 'Beneficiary':('Customer-owned destination record and configured transfer cap.',
   'BeneficiaryID determines the saved destination. IFSC_Code determines BankName only in Bank_Directory, preventing that transitive dependency here. LinkedAccountNo may be external and is not a local account foreign key.'),
 'Fixed_Deposit':('Customer investment principal, term, rate and generated maturity value.',
   'FD_ID determines the independently stored deposit facts. Customer details are separated. MaturityAmount is generated from Amount, InterestRate, StartDate and EndDate; see the formal derived-column qualification below.'),
 'Employee':('A branch official with role, unique email and employment details.',
   'EmployeeID and the unique Email are candidate keys. Salary and HireDate are employee-specific; Role does not determine a fixed salary, and branch descriptions are not repeated.'),
 'Audit_Log':('Independent chronological evidence of record-level database activity.',
   'LogID determines the event facts. TableName plus RecordID identifies a referenced business record but may appear in many audit events. Neither determines the event description or timestamp.'),
}

CHANGES=[
 ('Branch and Postal_Directory','Branch adds Address, Pincode and ContactPhone. Table 1 Postal_Directory stores City and State once per Pincode.','Removes Pincode -> City, State from Branch; avoids inconsistent edits to shared location facts.'),
 ('Beneficiary and Bank_Directory','Beneficiary adds IFSC_Code and MaxTransferLimit. Table 2 Bank_Directory stores IFSC_Code -> BankName.','Separates a destination-bank fact from a customer-specific saved payee and cap.'),
 ('Customer','Adds Gender, Age, Address and PAN_Number; Age >= 18 and allowed Gender CHECKs; PAN uniqueness.','Makes KYC attributes explicit and enforces adult ownership and identity uniqueness.'),
 ('Loan','Adds LoanType, TenureMonths, IssueDate, OutstandingBalance and generated stored EMI_Amount.','Database computes EMI consistently from the current contract inputs; the generated value cannot be independently edited.'),
 ('Fixed_Deposit','Adds StartDate, EndDate, Status and generated stored MaturityAmount.','EndDate must follow StartDate; the database enforces one maturity expression.'),
 ('Other operational entities','Transaction_History adds PaymentMode and Remarks; Loan_Repayment adds PaymentMode and PaymentStatus; Account adds InterestRate and OpeningDate; Employee adds Email, Salary and HireDate; Audit_Log adds TableName and RecordID.','Makes channels, lifecycle states, product terms and audit targets available for demonstrable workflows.'),
]

CATEGORIES=[
 ('A Basic','SELECT, WHERE, ORDER BY, DISTINCT, BETWEEN, IN, LIKE, IS NULL','6','6'),
 ('B Aggregate','COUNT, SUM, AVG, MIN, MAX, GROUP BY, HAVING; WHERE versus HAVING','5','5'),
 ('C Joins','INNER, LEFT, RIGHT, FULL OUTER, SELF and meaningful CROSS','6','6'),
 ('D Subqueries','Scalar, single-row, multi-row IN/ANY/ALL, correlated, EXISTS, NOT EXISTS, NOT IN with NULL fix','6','6'),
 ('E Set operations','UNION, UNION ALL, INTERSECT, EXCEPT','4','4'),
 ('F Advanced functions','CASE, COALESCE, NULLIF, AGE/EXTRACT, UPPER/CONCAT, ROUND','4-6','4'),
 ('G Window and analytic','RANK/DENSE_RANK, ROW_NUMBER, running SUM, LAG/LEAD','4','4'),
]

EXPECTED={
 'CQ-01':'One row per Defaulted loan, with the borrower, issuing branch and outstanding principal selected once. Failed-attempt count is zero when no failed attempt exists; a missing successful payment yields NULL. The result excludes all other loan statuses.',
 'CQ-02':'One activity total per branch is ranked by recorded events in the last 365 days. Up to the top three distinct ranks are returned, with ties retained and zero-activity branches preserved before ranking.',
 'CQ-03':'An Approved customer qualifies only when both an Active FD and an Active account exist and no Loan of any status exists. Separate product aggregates prevent an FD amount from being multiplied by account count.',
 'CQ-04':'One row per non-NULL stored transaction date is ranked by gross event amount. Only the three highest distinct volume ranks remain. Two history events for one internal transfer contribute twice to gross volume, not net funds moved.',
 'CQ-05':'Every customer remains visible. Missing product totals become zero; indicative net position equals non-Closed account balances plus Active FD principal minus issued loan outstanding. Pending and Closed loans are excluded here.',
 'CQ-06':'Only Active FDs whose EndDate is today or earlier appear. The oldest Active account supplies a suggested contact branch; a customer without an Active account still has a due FD row with NULL contact fields.',
 'CQ-07':'All branches are shown once with independently aggregated payroll and deposits. Customers are distinct within the branch. Ratios become NULL where staff count or payroll is zero.',
 'CQ-08':'There is one customer and destination-bank group, ranked within that bank by the sum of configured per-beneficiary caps. This is configured exposure, not executed transfer value or a customer-wide enforcement rule.',
 'CQ-09':'Each loan is preserved even with no repayments. Only Success amounts enter collections; Pending and Failed counts remain separate. Principal and outstanding balances are never summed across installment rows.',
 'CQ-10':'Accounts with unapproved or NULL KYC, Dormant status, or Frozen status appear with historical event totals. At most one latest Account audit is attached through both TableName and RecordID. Historical activity alone does not prove a post-freeze breach.',
}

OBJECT_PURPOSE={
 'sp_transfer_funds':'Validate two different Active accounts, a finite positive paise-accurate amount, an allowed payment mode and sufficient funds. Lock accounts in ascending AccountNo order, debit and credit once, write two history events, and explicitly COMMIT.',
 'sp_pay_loan_emi':'Lock the owned Active account and the repayable owned loan, deduct LEAST(EMI_Amount, OutstandingBalance), log a withdrawal and insert a successful repayment. The repayment trigger changes debt and closes the final installment.',
 'fn_calculate_emi':'Return the rounded monthly compound-interest installment using P*r*(1+r)^n / ((1+r)^n-1), r=annual percent/1200. Reject invalid finite inputs; extend the function to zero interest as P/n.',
 'fn_customer_net_worth':'Return account balances plus Active FD principal minus all stored loan outstanding for an existing customer. Independent aggregates and COALESCE prevent join fan-out and NULL arithmetic.',
 'trg_audit_account_changes':'AFTER UPDATE OF Balance on Account, write exactly one old/new balance audit when the values are distinct. A no-op update creates no audit; a rolled-back update also rolls back its audit.',
 'trg_loan_repayment':'AFTER INSERT on Loan_Repayment, act only on Success. Lock the loan FOR NO KEY UPDATE, reject overpayment, reduce outstanding, and change Status to Closed when the balance reaches zero.',
 'sp_batch_credit_savings_interest':'Use a bound explicit savings_cursor FOR loop over Active Savings accounts under row locks. Post rounded Balance*InterestRate/1200 as Deposit/Internal, with Monthly Interest Credit remarks. Skip accounts already posted in the current calendar month.',
 'sp_batch_mature_fixed_deposits':'Use a bound explicit maturity_cursor FOR loop for Active FDs due by CURRENT_DATE. Credit the oldest Active Savings account, mark the FD Matured, record history and an FD_MATURED audit. Abort the entire batch if an eligible payout account is absent.',
}

CUSTOMER_MODULES=[
 ('Portfolio Overview','Customer identity selected at login; Refresh portfolio.','GET /api/customer/{id}/overview; fn_customer_net_worth(id); independent Account, Fixed_Deposit and Loan aggregates.','Net worth, liquid balance, Active FD principal and maturity, loan liability, KYC and beneficiary count; asset/liability chart and recent transaction feed.'),
 ('My Accounts and Passbook','Account cards; Deposit and Withdraw dialogs with owned AccountNo, amount, mode and remarks; statement filters for account, type and channel.','GET /api/customer/{id}/accounts and /transactions; POST /api/transaction/deposit or /withdraw; SELECT FOR UPDATE, Account UPDATE, Transaction_History INSERT, audit trigger.','Exact account ownership, servicing branch and IFSC, rates, status, opening dates and balances. Successful writes refresh the statement; invalid or Frozen/Dormant accounts show rejection.'),
 ('Fund Transfers and Beneficiaries','Choose internal account or saved beneficiary; select sender, receiver/payee, amount, mode, remarks. Add payee name, external-capable account, existing/new IFSC and limit; delete a saved payee.','POST /api/transaction/transfer; internal CALL sp_transfer_funds; saved-payee lock and _bank_transfer_atomic for internal settlement, or an atomic outward debit for external destinations. GET/POST /api/customer/{id}/beneficiaries; DELETE /api/beneficiaries/{id}?customer_id=... .','Commit/error notification and refreshed balances/history. Daily limits sum tagged outgoing events across owned sender accounts. External payments are recorded local outward events, not real payment-rail settlement.'),
 ('Loans and EMI Center','Select a loan; Pay Monthly EMI with owned funding AccountNo; View repayments.','GET /api/customer/{id}/loans; GET /api/loans/{id}/repayments?customer_id=...; POST /api/loans/pay-emi -> CALL sp_pay_loan_emi.','Principal, rate, tenure, generated EMI, outstanding, status and paid-principal progress; real Loan_Repayment history refreshes after the trigger fires. No fictional amortization or due-date schedule is generated.'),
 ('Fixed Deposit Vault','Open New Fixed Deposit: funding AccountNo, principal >= INR 1,000, annual rate, StartDate and EndDate; live estimated maturity.','GET/POST /api/customer/{id}/fixed-deposits; one transaction debits the owned Active account, logs funding and INSERTs Fixed_Deposit without MaturityAmount.','FD lifecycle table and database-generated maturity; a failed date CHECK rolls back both the FD insert and prior funding debit.'),
 ('Profile and KYC Details','Select the customer using the global identity switcher; read-only KYC record.','GET /api/customer/{id}/overview -> Customer.','CustomerID, full name, Gender, Age, Address, Phone, Email, PAN_Number and KYC_Status. KYC approval is an official workflow.'),
]

ADMIN_MODULES=[
 ('Executive Analytics and Views','Refresh records; bank-wide KPI, chart and table inspection.','GET /api/admin/dashboard and /summary; vw_branch_financial_summary; vw_active_vellore_accounts; independent product aggregates and all 12 table counts.','CASA deposits, active loan and FD books, customers/branches/staff, branch deposits versus loans, account/loan distributions, payment-mode volume, and live view results.'),
 ('Customer 360 and KYC Directory','Search by identity/PAN/account; KYC filter; Inspect 360; Approve/Reject; Onboard New Customer with KYC/contact fields.','GET /api/admin/customers-master?search=&kyc_status=; GET /api/admin/customer/{id}/full-360; POST /api/admin/customers; PATCH /api/admin/customers/{id}/kyc.','Owner-to-account chips, total balances, active loan/FD counts, complete portfolio modal and KYC feedback. Advisory-lock ID allocation and database UNIQUE/CHECK constraints protect onboarding.'),
 ('Account Management Desk','Search/filter master accounts; open account by Customer, Branch, Savings/Current/Salary, deposit and rate; set Active, Dormant, Frozen or Closed.','GET/POST /api/admin/accounts; PATCH /api/admin/accounts/{account_no}/status; opening deposit is recorded in Transaction_History in the same transaction.','Owner identity, account number, branch, type, balance and rate; immediate status updates. A frozen account is subsequently rejected by the transfer/withdrawal logic.'),
 ('Loan Origination and Portfolio','Borrower, issuing branch, Home/Personal/Auto/Education/Business, principal, rate, months; optional borrower disbursement account.','GET/POST /api/admin/loans; GET /api/admin/emi-preview -> fn_calculate_emi; generated Loan.EMI_Amount; optional Account credit/history in one transaction.','Live database EMI preview, loan master with ownership and terms, Approved loan without disbursement or Active loan with disbursement. Only the borrower accounts appear as funding destinations.'),
 ('Fixed Deposit Book','Status filter Active/Matured/Closed; search across customer investments.','GET /api/admin/fixed-deposits?status= -> Fixed_Deposit JOIN Customer.','All FD IDs, owner names, principal, rates, dates, generated maturity and lifecycle state.'),
 ('Branches Staff and Lookup Directories','Sub-tabs for branch network, officials, postal and IFSC directories; add branch/pincode, bank IFSC or employee using normalized lookups.','GET /api/admin/infrastructure; POST /api/admin/branches, /employees, /postal-directory and /bank-directory.','Branch locations/contacts; employee roster; canonical Pincode -> City/State and IFSC -> BankName. Conflicting existing shared facts are rejected rather than silently overwritten.'),
 ('Batch Cursors and Live Trigger Audit','Run savings interest or FD maturity batch; filter audit by TableName and ActionType; Refresh records.','POST /api/admin/batch/interest and /mature-fds -> corresponding cursor procedures; GET /api/admin/audit-log and /programmability.','Processed counts, latest 100 audit rows polled every 10 seconds while the tab is visible, and installed view/routine/trigger catalogue. Duplicate posting guards make immediate batch reruns no-ops.'),
]


def shade(element, color):
    prop=element.get_or_add_tcPr() if element.tag==qn('w:tc') else element.get_or_add_pPr()
    shd=prop.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); prop.append(shd)
    shd.set(qn('w:fill'),color)


def run_style(run, size=11, bold=False, font='Times New Roman', color='000000'):
    run.font.name=font; run.font.size=Pt(size); run.bold=bold; run.font.color.rgb=RGBColor.from_string(color)


def text_value(value):
    if value is None: return 'NULL'
    if isinstance(value,bool): return 'true' if value else 'false'
    return str(value)


class Report:
    def __init__(self, path, evidence, review, page_map=None):
        self.d=Document(path); self.e=evidence; self.review=review; self.toc=[]; self.page_map=page_map or {}
        self.screenshots={}; self.ui_screenshots={}
        if review in (2,3):
            folder=ROOT/'report_screenshots'/'review3'
            if (folder/'sql_manifest.json').exists():
                self.screenshots={s['id']:s for s in json.loads((folder/'sql_manifest.json').read_text(encoding='utf-8'))}
            if review==3 and (folder/'ui_manifest.json').exists():
                self.ui_screenshots={s['name']:s for s in json.loads((folder/'ui_manifest.json').read_text(encoding='utf-8'))}
        body=self.d.element.body
        # Retain the exact cover and drawing through the VIT logo paragraph.
        original_children=list(body)
        self.grading=deepcopy(self.d.tables[-1]._tbl)
        self.original_grading_xml=self.d.tables[-1]._tbl.xml
        for child in original_children[21:]:
            if child.tag!=qn('w:sectPr'): body.remove(child)
        self.configure_styles()
        self.cover()

    def configure_styles(self):
        # Body typography follows the template's Times New Roman baseline.
        for name in ['Normal','Body Text']:
            style=self.d.styles[name]; style.font.name='Times New Roman'; style.font.size=Pt(11)
            style.font.color.rgb=RGBColor(0,0,0)
            style.paragraph_format.line_spacing=1.08
            style.paragraph_format.space_after=Pt(5)
            style.paragraph_format.space_before=Pt(0)
        for name,size in [('Heading 1',14),('Heading 2',12),('Heading 3',11),('Title',18)]:
            s=self.d.styles[name]; s.font.name='Times New Roman'; s.font.size=Pt(size)
            s.font.color.rgb=RGBColor(0,0,0); s.font.bold=True
            s.paragraph_format.keep_with_next=True
            s.paragraph_format.space_before=Pt(12); s.paragraph_format.space_after=Pt(5)
            pp=s.element.get_or_add_pPr()
            for tag in ['numPr','shd','pBdr']:
                node=pp.find(qn('w:'+tag))
                if node is not None: pp.remove(node)
        for name,font,size in [('SQL Code','Consolas',9),('Result Code','Consolas',9),('Report TOC','Times New Roman',11)]:
            if name not in self.d.styles: self.d.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
            s=self.d.styles[name]; s.font.name=font; s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0)
            s.paragraph_format.space_before=Pt(0); s.paragraph_format.space_after=Pt(0)
            s.paragraph_format.line_spacing=1.0
            s.paragraph_format.keep_with_next=False
            if name=='SQL Code':
                s.paragraph_format.left_indent=Inches(.08); s.paragraph_format.right_indent=Inches(.08)
                shade(s.element,'F2F4F7')
        settings=self.d.settings.element
        update=OxmlElement('w:updateFields'); update.set(qn('w:val'),'true'); settings.append(update)

    def cover(self):
        p=self.d.paragraphs
        for paragraph in p:
            paragraph.paragraph_format.space_before=Pt(0)
            paragraph.paragraph_format.space_after=Pt(0)
        replacements={1:TITLE,3:f'PROJECT REPORT  REVIEW {self.review}',6:f'{MEMBERS[0][0]} ({MEMBERS[0][1]})',
                      7:f'{MEMBERS[1][0]} ({MEMBERS[1][1]})',8:f'{MEMBERS[2][0]} ({MEMBERS[2][1]})',
                      12:'DATABASE SYSTEMS (BCSE302P) -Lab PROJECT',16:'Dr. Anand Bihari'}
        for index,value in replacements.items():
            p[index].clear(); p[index].add_run(value)
            p[index].alignment=WD_ALIGN_PARAGRAPH.CENTER
            run_style(p[index].runs[0],18 if index==1 else 14,index in [6,7,8,12,16])
        p[1].style='Title'; p[1].paragraph_format.line_spacing=1
        p[1].paragraph_format.space_before=Pt(0); p[1].paragraph_format.space_after=Pt(0)
        for index in [6,7,8]: p[index].paragraph_format.space_after=Pt(0)

    def para(self,text='',bold=False,size=None,parent=None):
        shot=re.match(r'\[Screenshot Placeholder: ([\w-]+)',str(text))
        if shot and shot[1] in self.screenshots:
            return self.image(shot[1],parent=parent)
        p=(parent or self.d).add_paragraph()
        p.paragraph_format.keep_with_next=False
        p.paragraph_format.widow_control=True
        r=p.add_run(str(text)); run_style(r,size or 11,bold)
        return p

    def heading(self,text,level=1,newpage=False,toc=True):
        p=self.d.add_paragraph(text,style=f'Heading {level}')
        if newpage: p.paragraph_format.page_break_before=True
        if toc: self.toc.append((text,level))
        return p

    def code(self,text,parent=None,result=False):
        holder=parent or self.d
        lines=text.strip().splitlines() or ['']
        p=holder.add_paragraph(style='Result Code' if result else 'SQL Code')
        p.paragraph_format.keep_with_next=False
        p.paragraph_format.keep_together=len(lines)<=14
        p.paragraph_format.widow_control=False
        run=p.add_run('\n'.join(line.expandtabs(4) for line in lines))
        run_style(run,9,font='Consolas')
        if not result: shade(p._p,'F2F4F7')
        self.para('',size=2,parent=holder).paragraph_format.space_after=Pt(2)

    def table(self,headers,rows,widths=None,parent=None,font=9.5):
        holder=parent or self.d
        t=holder.add_table(rows=1,cols=len(headers)) if parent is None else holder.add_table(rows=1,cols=len(headers))
        t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
        available=7 if parent is None else 5.45
        widths=widths or [available/len(headers)]*len(headers)
        for col,width in zip(t.columns,widths): col.width=Inches(width)
        for cell,value in zip(t.rows[0].cells,headers): cell.text=str(value)
        for values in rows:
            for cell,value in zip(t.add_row().cells,values): cell.text=text_value(value)
        self.style_table(t,widths,font,header=True)
        self.para('',size=2,parent=holder).paragraph_format.space_after=Pt(3)
        return t

    def style_table(self,t,widths,font=9.5,header=True):
        props=t._tbl.tblPr
        borders=OxmlElement('w:tblBorders')
        for tag in ['top','left','bottom','right','insideH','insideV']:
            e=OxmlElement('w:'+tag); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
        props.append(borders)
        for row_index,row in enumerate(t.rows):
            if header:
                no_split=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
            for cell,width in zip(row.cells,widths):
                cell.width=Inches(width); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                mar=OxmlElement('w:tcMar')
                for tag,n in [('top',65),('bottom',65),('left',90),('right',90)]:
                    e=OxmlElement('w:'+tag); e.set(qn('w:w'),str(n)); e.set(qn('w:type'),'dxa'); mar.append(e)
                cell._tc.get_or_add_tcPr().append(mar)
                if row_index==0 and header: shade(cell._tc,'DCE6F1')
                elif row_index%2: shade(cell._tc,'F7F9FC')
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(2); p.paragraph_format.space_before=Pt(0)
                    p.paragraph_format.line_spacing=1.0; p.paragraph_format.keep_with_next=bool(header and row_index==0)
                    for r in p.runs: run_style(r,font,row_index==0 and header)
            if row_index==0 and header:
                repeat=OxmlElement('w:tblHeader'); row._tr.get_or_add_trPr().append(repeat)

    def placeholder(self,text,parent=None):
        if text.startswith('[PLACEHOLDER: Insert UI Screenshot - '):
            name=text.removeprefix('[PLACEHOLDER: Insert UI Screenshot - ').removesuffix(']')
            if name in self.ui_screenshots:
                return self.image_record(self.ui_screenshots[name],name,parent)
        if 'DBeaver CREATE TABLE' in text and 'DDL' in self.screenshots:
            self.para('Existing-table verification in DBeaver. All twelve installed relations and their columns are queried from PostgreSQL metadata. This screenshot verifies the existing schema; the clean CREATE TABLE execution is documented by the separate disposable-database output above.',parent=parent,size=9.5)
            return self.image('DDL',parent=parent)
        p=self.para(text,size=10,parent=parent)
        p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(7)
        shade(p._p,'FFF2CC')
        for r in p.runs: r.bold=True
        return p

    def image_record(self,record,label,parent=None):
        p=(parent or self.d).add_paragraph()
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_together=True
        p.paragraph_format.keep_with_next=True
        p.paragraph_format.space_before=Pt(4)
        p.paragraph_format.space_after=Pt(2)
        p.add_run().add_picture(str(Path(record['path'])),width=Inches(5.45 if parent is not None else 6.9))
        self.para(f'{label} — live screenshot, {record["captured_at"][:10]}.',size=9,parent=parent)
        return p

    def image(self,key,parent=None):
        record=self.screenshots.get(key)
        if record:
            return self.image_record(record,key,parent)

    def view_output(self,name,parent):
        self.output(self.e['views'][name],parent)
        self.image(name,parent)

    def diagram(self,path,caption):
        from PIL import Image
        with Image.open(path) as im:
            width,height=im.size
        scale=min(6.9/width,6.1/height)
        p=self.d.add_paragraph()
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next=True
        p.paragraph_format.keep_together=True
        p.add_run().add_picture(str(path),width=Inches(width*scale))
        self.para(caption,size=9)

    def output(self,result,parent=None,caption=True,sample=None):
        rows=result.get('rows',[]); columns=result.get('columns',[])
        if not columns:
            self.para(result.get('status','Command completed.'),size=10,parent=parent); return
        count=result.get('row_count',len(rows))
        limit=sample if sample is not None else 1 if len(columns)>8 else 2 if len(columns)>6 else 3
        rows=rows[:limit]
        if caption: self.para(f'Actual output: {count} row(s); first {len(rows)} shown. NULL denotes a missing SQL value.',size=9.5,parent=parent)
        if not rows:
            self.para('0 rows returned. The query executed successfully; no record satisfies its predicate.',size=10,parent=parent); return
        # Vertical records are readable when a result is wider than the page.
        if len(columns)>6 or parent is not None:
            for i,row in enumerate(rows,1):
                block=f'Row {i}\n'+'\n'.join(f'{col}: {text_value(value)}' for col,value in zip(columns,row))
                self.code(block,parent,result=True)
        else:
            self.table([c.replace('_',' ') for c in columns],rows,parent=parent,font=9)

    def field_table(self,fields):
        t=self.d.add_table(rows=len(fields),cols=2); t.autofit=False
        self.style_table(t,[1.15,5.85],font=10,header=False)
        for row,(name,content) in zip(t.rows,fields):
            label,value=row.cells; label.text=name; shade(label._tc,'DCE6F1')
            for r in label.paragraphs[0].runs: run_style(r,10,True)
            value.text=''
            if callable(content): content(value)
            else: self.para(content,parent=value,size=10)
            # A default empty paragraph is required by Word but need not consume a line.
            first=value.paragraphs[0]
            if not first.text and len(value.paragraphs)>1: value._tc.remove(first._p)
        self.para('',size=2)
        return t

    def contents(self):
        self.heading('Table of Contents',level=1,newpage=True,toc=False)
        self.toc_anchor=self.d.add_paragraph()
        self.para('Read Sections 1 to 8 for the shared database implementation. Review 3 extends this foundation with final verification, concurrency, a timed demonstration and application modules. The evaluation grid is reserved for the faculty.',size=10)
        self.para(f'Evidence captured {self.e["captured_at"]}. Live row counts and query samples are a point-in-time snapshot; later UI transactions can change them.',size=10)

    def fill_contents(self):
        entries=[(x,l) for x,l in self.toc if l==1 or (l==2 and x.startswith(('8.','10.')))]
        for text,level in entries:
            p=self.d.add_paragraph(style='Report TOC')
            p.paragraph_format.space_after=Pt(4)
            p.paragraph_format.left_indent=Inches(.15 if level==2 else 0)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(6.8),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
            page=self.page_map.get(str(self.review),{}).get(text,'--')
            p.add_run(f'{text}\t{page}')
            self.toc_anchor._p.addprevious(p._p)
        self.toc_anchor._p.getparent().remove(self.toc_anchor._p)


def abstract_and_scope(r):
    r.heading('Abstract',newpage=True,toc=False)
    r.para('The Core Banking & Transaction Management System implements a relational banking laboratory in PostgreSQL with twelve connected operational entities. It supports KYC onboarding, multi-branch customer accounts, positive-valued transaction events, loan origination and repayment, fixed deposits, saved beneficiaries, employee records and independent audit events. Postal and bank directories isolate shared facts; primary, unique, foreign-key, domain and date constraints enforce the declared data model.')
    r.para('The database layer combines thirty-five standard SQL demonstrations and ten managerial business queries with three views, two business functions, two core procedures, two triggers and two explicit cursor batches. Fund transfers lock both accounts before posting and commit account, history and trigger effects together. A rejected validation or later SQL failure undoes the enclosing transaction. Successful loan repayments reduce outstanding debt through a trigger; stored generated expressions compute EMI and deposit maturity from their input terms.')
    r.para('A Python FastAPI service and React 19 interface expose customer and bank-official workflows. Six customer sections show exact account ownership, statements, transfers, beneficiaries, loans, deposits and KYC. Seven official sections combine analytics, customer 360-degree inspection, account controls, loan sanctioning, fixed-deposit books, infrastructure and batch audit evidence. Both light and dark display modes are supported.')
    r.para('Review 2 establishes the schema, SQL coverage, programmability, transactions, indexing and role grants. Review 3 retains that evidence and adds final test outcomes, an observed two-session lock conflict, a timed demonstration runbook and complete application input/output descriptions. The implementation is a university demonstration with simulated portal identity selection; database roles are verified separately. Historical fixtures represent illustrative events and opening balances, not a complete double-entry ledger.')
    r.contents()
    r.heading('1 Introduction',newpage=True)
    r.para('Banks maintain interdependent facts about customers, servicing branches and financial products. Repeating those facts in one large relation creates update anomalies and makes ownership, debt and audit history difficult to reconcile. Our project separates those responsibilities into twelve relations and executes financial mutations within database transaction boundaries.')
    r.para('PostgreSQL is the authority for declared keys, relationships, generated amounts, triggers and stored workflow validation. FastAPI uses parameterized Psycopg2 statements and isolates failed requests through rollback and connection cleanup. React presents the same underlying relationships to customers and officials, making a constraint failure, successful commit or trigger-generated event visible during the viva.')
    r.para('The independently stored facts are designed for third normal form. Stored generated financial expressions are treated explicitly as constrained derived projections: they cannot be independently supplied or updated. Their formal functional-dependency qualification is discussed in Section 6 rather than hidden by the interface.')
    r.heading('2 Project Scope')
    r.table(['Area','Implemented scope'],[
        ['Customer lifecycle','Adult KYC record, unique identity/contact fields, approval state, searchable customer discovery and exact account ownership.'],
        ['Banking products','Savings, Current and Salary accounts; five loan purposes; generated EMI; customer FDs with generated actual-day maturity.'],
        ['Transactional workflows','Deposit, withdrawal, internal transfer, saved beneficiary transfer, EMI payment, funded FD creation and optional loan disbursement.'],
        ['Operations and evidence','Branch/staff directories, aggregate views, cursor batches, balance audits, role grants, indexes, rollback tests and two-session row-lock proof.'],
        ['System boundary','Local academic deployment. No real interbank rails, authenticated internet banking, regulatory NPA aging engine, joint-account model, general ledger, password store or repayment due-date calendar.'],
    ],[1.4,5.6])
    r.heading('3 Key Contacts and Stakeholders')
    r.table(['Name','Registration Number','Phone Number'],MEMBERS,[3.2,1.7,2.1])
    r.para('Prepared for DATABASE SYSTEMS (BCSE302P) -Lab PROJECT. Submitted to Dr. Anand Bihari, School of Computer Science and Engineering. The allocation in Section 9 records the team responsibilities specified for these reviews; the faculty assigns the individual evaluation marks.')
    r.heading('4 Project Resource Requirements')
    r.heading('4.1 Software Resource Requirements',2)
    r.table(['Resource','Requirement and project use'],[
        ['Operating system','Windows 11, macOS or Linux; this evidence run used the local Windows project workspace.'],
        ['DBMS','PostgreSQL 16 or later; live server reports PostgreSQL 18.6. PL/pgSQL, transactional DDL and stored generated columns are used.'],
        ['Containerization','Docker; banking-postgres exposes PostgreSQL on localhost:5432.'],
        ['Backend','Python 3.x (project requires >= 3.12), FastAPI, Psycopg2 and Pydantic; uv manages dependencies and execution.'],
        ['Frontend','React 19, Vite, Tailwind CSS v4 with its Vite plugin, Axios, React Router, Recharts and Lucide Icons.'],
        ['Database client','DBeaver for manual SQL execution, catalogue inspection and evaluator screenshots; psql for scripted execution.'],
        ['System design tools','Mermaid.js flowcharts rendered in the browser as Chen rectangles, diamonds and ovals; source diagrams and PNG exports are retained in docs/diagrams. Draw.io can also edit a manually recreated version.'],
        ['Report tooling','python-docx, live PostgreSQL catalogue/result capture, Word pagination and rendered-page quality checks.'],
    ],[1.4,5.6])
    r.heading('4.2 Deployment and Reproducibility',2)
    r.code('uv sync\nuv run uvicorn backend.main:app --reload\n# In a separate terminal\ncd frontend\nnpm install\nnpm run dev')
    r.para('FastAPI runs at http://localhost:8000 with interactive schemas at /docs. React runs at http://localhost:5173. The reports use .report_work/evidence.json, generated by collect_report_evidence.py. Report regeneration starts from immutable backups in report_backups/original, not from previously expanded reports.')
    r.heading('5 ER Diagram',newpage=True)
    folder=ROOT/'docs'/'diagrams'
    if (folder/'chen-overview.png').exists():
        r.para('The Chen model below represents all twelve entities and eleven declared foreign-key relationships in sql/01_ddl.sql. The overview shows the relationship structure; four following attribute panels show all 81 columns without crowding the overview.',size=10)
        r.para('Notation: rectangle = entity; diamond = relationship; oval = attribute; underlined label = primary key; dashed oval = generated derived value. 1:N labels give maximum cardinality. Each child foreign key is mandatory, while a parent can have zero or many children. FK suffixes in the attribute panels identify physical reference columns.',size=9.5)
        r.diagram(folder/'chen-overview.png','Figure 5.1  Chen relationship overview of the twelve-table banking schema')
        r.para('Audit_Log is independent: its TableName and RecordID identify an audit target without a declared FK. Branch.IFSC_Code does not reference Bank_Directory, and Beneficiary.LinkedAccountNo does not reference Account. External beneficiary destinations are allowed. Fixed-deposit branch attribution is a reporting rule through the customer account, rather than a declared FD-to-Branch relationship.',size=9.5)
        panels=json.loads((folder/'diagrams.json').read_text(encoding='utf-8'))
        for i,panel in enumerate(panels[1:],2):
            r.heading(f'5.{i-1} {panel["title"]}',2,newpage=True,toc=False)
            r.para(panel['note'],size=10)
            r.diagram(folder/(panel['id']+'.png'),f'Figure 5.{i}  {panel["title"]}')
        r.para('EMI_Amount is derived from PrincipalAmount, InterestRate and TenureMonths. MaturityAmount is derived from Amount, InterestRate, StartDate and EndDate. PostgreSQL stores both as GENERATED ALWAYS expressions; their exact formulas and normalization qualification are documented in Section 6. Editable Mermaid sources, SVGs and browser-exported PNGs are retained in docs/diagrams.',size=10)
    else:
        r.placeholder('[PLACEHOLDER: PASTE UPDATED 12-TABLE CHEN NOTATION ER DIAGRAM HERE]')


def schema_section(r):
    r.heading('6 Tables and Constraints',newpage=True)
    r.heading('6.1 Design Changes Since Review 1',2)
    r.para('The following change log records the requested Review 1 feedback baseline. Postal and bank lookup decomposition removes repeated directory facts; operational columns make constraints and workflow outcomes demonstrable. The current DDL is the authority for exact implementation details.')
    r.table(['Changed component','Upgrade','Reason and integrity effect'],CHANGES,[1.25,3.25,2.5])
    r.heading('6.2 Normalization and Derived Attributes',2)
    r.para('1NF: each relation has scalar attributes, one value per cell and a declared primary key; repeating products and payment attempts are separate rows. 2NF: every declared primary key is a single column, so no non-key fact depends on only part of a composite key. 3NF: owner, branch, postal and bank-name facts are separated according to the functional dependencies described for each relation below.')
    r.para('Formal qualification for generated attributes. In the independently stored Loan facts, LoanID determines the terms and state. However, (PrincipalAmount, InterestRate, TenureMonths) also determines EMI_Amount, and is not a superkey. Likewise (Amount, InterestRate, StartDate, EndDate) determines MaturityAmount in Fixed_Deposit. If those stored generated outputs are counted as ordinary relation attributes, strict textbook 3NF is not satisfied by that extension. The base business facts are normalized and the redundant outputs are controlled, immutable generated expressions. This is more precise than claiming that GENERATED ALWAYS automatically proves 3NF.')
    r.para('No additional normalization rule is silently enforced: CHECK constraints can allow NULL unless NOT NULL is also declared. Defaulted KYC/product statuses and some timestamps are nullable in the exact schema. Branch.IFSC_Code is unique but has no Bank_Directory FK; PAN format and interest-rate positivity outside Loan are principally application/domain decisions unless an explicit CHECK is listed.')
    ddl=r.e['scripts']['01_ddl.sql']
    creates={re.match(r'CREATE TABLE (\w+)',clean_sql(s))[1]:clean_sql(s) for s in split_sql(ddl) if clean_sql(s).startswith('CREATE TABLE ')}
    columns=r.e['columns']; constraints=r.e['constraints']
    r.heading('6.3 Table Definitions and Creation Scripts',2)
    for n,name in enumerate(TABLES,1):
        desc,norm=TABLE_INFO[name]
        r.heading(f'6.3.{n} {name}',3)
        r.para(f'TABLE NAME: {name}',bold=True).paragraph_format.keep_with_next=True
        r.para(f'DESCRIPTION: {desc}').paragraph_format.keep_with_next=True
        form='3NF base facts with a controlled stored generated extension' if name in ['Loan','Fixed_Deposit'] else '3NF'
        r.para(f'NORMAL FORM: {form}. 1NF uses atomic values. 2NF follows from the single-column primary key. 3NF justification: {norm}').paragraph_format.keep_with_next=True
        rows=[]
        for col in columns['rows']:
            values=dict(zip(columns['columns'],col))
            if values['table_name']!=name.lower(): continue
            colname=values['column_name']; datatype=values['data_type']
            if datatype=='character varying': datatype=f"VARCHAR({values['character_maximum_length']})"
            elif datatype=='numeric': datatype=f"DECIMAL({values['numeric_precision']},{values['numeric_scale']})"
            elif datatype=='timestamp without time zone': datatype='TIMESTAMP'
            elif datatype=='integer': datatype='INT'
            else: datatype=datatype.upper()
            notes=['NOT NULL' if values['is_nullable']=='NO' else 'NULL allowed']
            if values['column_default']:
                default=values['column_default']
                if 'nextval(' in default: datatype='SERIAL'; notes.append('sequence default')
                else: notes.append('DEFAULT '+default)
            if values['is_generated']=='ALWAYS': notes.append('GENERATED ALWAYS AS (...) STORED; full expression in DDL below')
            for record in constraints['rows']:
                tab,cname,kind,definition=record
                if tab==name.lower() and re.search(r'\b'+re.escape(colname)+r'\b',definition):
                    notes.append(definition)
            rows.append([colname,datatype,'; '.join(dict.fromkeys(notes))])
        r.table(['Attribute','Data Type','Constraint'],rows,[1.4,1.35,4.25],font=9.5)
        r.para('CREATE TABLE source from sql/01_ddl.sql',bold=True,size=10)
        r.code(creates[name])
    r.heading('6.4 Table Relationships',2)
    relationships=[]
    for table,con,definition in r.e['foreign_keys']['rows']:
        child=re.search(r'FOREIGN KEY \((.*?)\)',definition)[1]
        target=re.search(r'REFERENCES (\w+)\((.*?)\)',definition)
        action='CASCADE' if 'ON DELETE CASCADE' in definition else 'NO ACTION (default)'
        relationships.append([target[1],table,child,f'1:N; {action}'])
    r.table(['Parent','Child','Child foreign key','Cardinality and delete rule'],relationships,[1.5,1.7,1.5,2.3])
    r.para('All eleven foreign keys are NOT NULL on the child. Customer deletion cascades only to Account and Beneficiary; Loan deletion cascades to Loan_Repayment. Other references use PostgreSQL NO ACTION and can still block a cascade when protected history exists. Audit_Log.RecordID is a polymorphic identifier, so the audit table has no FK; queries must combine TableName with RecordID to avoid matching unrelated entities.')
    r.heading('6.5 Live Row Counts and DDL Verification',2)
    r.table(['Table','Rows in banking_db'],r.e['counts'],[4.8,2.2])
    r.para('The live snapshot includes the initial seed, bulk fixtures and subsequent demonstration activity. seed_bulk_data.py adds 100 Customer, 150 Account, 50 Loan, 500 Transaction_History and 40 Fixed_Deposit rows per successful run using Faker en_IN and execute_values in one transaction. Never rerun it merely to reproduce an existing count.')
    r.para('Clean DDL and DML execution was verified separately in a disposable PostgreSQL database using the unchanged sql/01_ddl.sql and sql/02_dml.sql. Both transactions committed; the advanced verification blocks also passed. The live banking_db was not reset.')
    r.code(r.e['clean_install']['verification_output'],result=True)
    r.placeholder('[PLACEHOLDER: Insert DBeaver CREATE TABLE Execution Screenshots]')
    r.heading('6.6 DML Operations',2)
    r.code("BEGIN;\nINSERT INTO Beneficiary\n(BeneficiaryID,CustomerID,BeneficiaryName,LinkedAccountNo,IFSC_Code)\nVALUES (-1,1,'Temporary DML demonstration',999999999999,'NMBK0000001');\nUPDATE Beneficiary SET MaxTransferLimit=25000 WHERE BeneficiaryID=-1;\nDELETE FROM Beneficiary WHERE BeneficiaryID=-1;\nCOMMIT;")
    r.para('The exact INSERT -> UPDATE -> DELETE sequence appears in sql/02_dml.sql and was executed during the clean verification. Its temporary negative ID leaves no final beneficiary row. The twelve live counts above document the operational dataset rather than the smaller rebuild fixture.')


def standard_section(r):
    r.heading('7 SQL Implementation',newpage=True)
    r.table(['Category','Required coverage','Min queries'],[[a,b,f'{c}; implemented {d}'] for a,b,c,d in CATEGORIES],[1.25,4.3,1.45])
    r.para('Exactly 35 executable statements are reproduced from sql/03_standard_queries.sql: A=6, B=5, C=6, D=6, E=4, F=4 and G=4. Category F meets the required range of four to six queries, with all required advanced functions demonstrated across four statements.')
    r.para('All results below were executed against banking_db in one repeatable-read, read-only snapshot. Small result samples preserve query order; the total result count is stated. Wide results use column/value records to remain legible. Monetary values are INR. Transaction sums describe gross positive event volume, not a reconstructed account balance.')
    last=None
    for q in r.e['standard']:
        cat=q['id'][0]
        if cat!=last:
            title=next(a for a,b,c,d in CATEGORIES if a.startswith(cat))
            r.heading(f'Category {title}',2); last=cat
        r.heading(f"{q['id']} {q['label'].rstrip('.')}",3)
        if q['id']=='B-04':
            p=r.para('WHERE filters individual rows before GROUP BY. HAVING filters the aggregate branch groups afterward. Removing non-Active accounts with HAVING alone would let their balances enter the sum; applying SUM(Balance) in WHERE is invalid at that query level.',bold=True)
            shade(p._p,'FFF2CC')
        if q['id']=='D-06':
            p=r.para('NULL pitfall: an unmatched value NOT IN a set containing NULL evaluates to UNKNOWN and is discarded by WHERE. Filtering inner NULLs restores the intended anti-membership test. The single statement prints the unsafe and corrected counts side by side.',bold=True)
            shade(p._p,'FFF2CC')
        r.code(q['sql'])
        r.output(q['result'])
        r.para(f'[Screenshot Placeholder: {q["id"]} SQL and result in DBeaver]',size=9)


def advanced_section(r):
    r.heading('8 Advanced SQL Implementation',newpage=True)
    r.heading('8.1 Complex Business Queries',2)
    r.para('Each CQ statement uses at least three base relations directly or through a multi-table CTE, with independent aggregation or analytic constructs where needed. The expected-output row states a hand-reasoned relational prediction; the actual-output row is the live result. Empty eligibility results are retained as valid evidence rather than filled with invented rows.')
    for q in r.e['complex']:
        r.heading(q['id'],3)
        def actual(cell,q=q):
            r.output(q['result'],cell)
            r.para(f'[Screenshot Placeholder: {q["id"]} execution]',parent=cell,size=9)
        r.field_table([
            ('ID',q['id']),('Business question',q['business_question']),
            ('SQL',lambda cell,q=q:r.code(q['sql'],cell)),
            ('Expected output',EXPECTED[q['id']]),('Actual output screenshot',actual),
            ('Explanation',q['explanation']),
        ])
    r.heading('8.2 Database Views',2)
    views={
      'vw_customer_passbook':('Readable passbook joining Customer, Account, Branch and Transaction_History. COALESCE replaces missing narrative only for display.',
          'Not automatically updatable. Its top-level FROM contains four joined relations and includes expression columns; PostgreSQL cannot unambiguously route arbitrary view writes to one base table.'),
      'vw_branch_financial_summary':('One row per branch, independently aggregated Active account count, stored loan outstanding and Active FD maturity attribution.',
          'Not automatically updatable. GROUP BY aggregates, CTEs and multiple joined sources produce calculated summaries rather than a single directly writable relation.'),
      'vw_active_vellore_accounts':('Restrict visible Account rows to Active accounts serviced by a branch whose postal city is Vellore. security_barrier and CASCADED CHECK OPTION protect the predicate boundary.',
          'Automatically updatable: the top-level FROM has only Account and its output columns are direct references. Branch and Postal_Directory are consulted through an EXISTS semi-join. CASCADED CHECK OPTION rejects writes that leave the view. Catalogue is_updatable=YES; the view alone is not a complete access-control policy.'),
    }
    for name,(purpose,useful) in views.items():
        r.heading(name,3)
        order='BranchID' if 'summary' in name else 'AccountNo'
        r.field_table([
          ('Purpose',purpose),
          ('CREATE VIEW statement',lambda cell,name=name:r.code(r.e['objects'][name]['sql'],cell)),
          ('Query using the view',lambda cell,name=name,order=order:r.code(f'SELECT * FROM {name} ORDER BY {order};',cell)),
          ('Screenshot / Live Output',lambda cell,name=name:r.view_output(name,cell)),
          ('Why useful',useful),
        ])
    r.para('Actual PostgreSQL updatability catalogue',bold=True)
    r.output(r.e['view_updatability'])

    groups=[('8.3 Stored Procedures',['sp_transfer_funds','sp_pay_loan_emi']),
            ('8.4 User Defined Functions',['fn_calculate_emi','fn_customer_net_worth']),
            ('8.5 Triggers',['trg_audit_account_changes','trg_loan_repayment']),
            ('8.6 Explicit Cursors',['sp_batch_credit_savings_interest','sp_batch_mature_fixed_deposits'])]
    for title,names in groups:
        r.heading(title,2)
        for name in names:
            r.heading(name,3); r.para('Purpose: '+OBJECT_PURPOSE[name])
            if name.startswith('trg_'):
                support='tg_audit_account_changes' if 'audit' in name else 'tg_loan_repayment'
                r.code(r.e['objects'][support]['sql'])
            r.code(r.e['objects'][name]['sql'])
            if name.startswith('fn_'):
                command='SELECT fn_calculate_emi(2500000,8.50,240), fn_calculate_emi(12000,0,12);' if 'emi' in name else 'SELECT fn_customer_net_worth(1);'
                r.para('Test execution against banking_db',bold=True,size=10); r.code(command)
                udf=r.e['udfs']; indices=[0,1] if 'emi' in name else [2]
                r.output({'columns':[udf['columns'][i] for i in indices], 'rows':[[udf['rows'][0][i] for i in indices]],'row_count':1})
            else:
                live=r.e['objects'][name].get('live_test')
                if live:
                    r.para('Live banking_db verification',bold=True,size=10)
                    r.code(live['command']); r.para(live['context']+'. Observed '+r.e.get('live_object_captured_at',r.e['captured_at'])+'. This object demonstration is later than the shared query snapshot; intervening presenter activity can change balances.',size=9.5)
                    r.output(live['output'])
                test=r.e['objects'][name]['test']; r.para('Test execution',bold=True,size=10); r.code(test['command'])
                r.para('Execution context: '+test['context']+'. These outputs are isolated commit/trigger demonstrations; live read-only evidence elsewhere remains from banking_db.',size=9.5)
                if 'before_count' in test: r.para(f'Audit count before {test["before_count"]}; after update {test["after_count"]}; both update and audit then rolled back.',size=10)
                if 'before_outstanding' in test: r.para(f'Outstanding before insert: INR {test["before_outstanding"]}. The successful test repayment is INR 10.',size=10)
                r.output(test['output'])
            r.image('TC-10' if name=='sp_transfer_funds' else name)
            benefits={
              'sp_transfer_funds':'Atomicity includes both account updates, two histories and audit trigger inserts. The caller must use top-level autocommit CALL because this procedure explicitly COMMITs; a validation exception occurs before that COMMIT. Both rows are locked in deterministic order to reduce opposite-direction deadlocks.',
              'sp_pay_loan_emi':'The account debit, withdrawal event, repayment row and triggered debt change share one transaction. The final payment is capped at outstanding. The lab deliberately deducts the entire installment from principal; it has no interest/principal amortization allocation table.',
              'fn_calculate_emi':'IMMUTABLE is appropriate because the result depends only on inputs. Loan requires positive interest, but the UDF also handles the mathematically valid zero-rate test. Stored Loan.EMI_Amount is checked against the function in TC-13.',
              'fn_customer_net_worth':'STABLE reads a consistent statement snapshot. Missing product rows contribute zero. Unlike CQ-05, this exact function includes every stored loan status and every account balance; the report does not silently equate those distinct metrics.',
              'trg_audit_account_changes':'The WHEN predicate suppresses meaningless no-op records. Audit and balance changes are transactional together. SQL script 06 makes this support function SECURITY DEFINER with a fixed search_path so permitted manager account updates can still write the protected log.',
              'trg_loan_repayment':'Only successful attempts affect debt. FOR NO KEY UPDATE serializes debt changes without conflicting with another repayment insert\'s foreign-key key-share lock. Overpayments are rejected instead of being silently discarded.',
              'sp_batch_credit_savings_interest':'The explicit cursor opens, fetches and closes through a bound FOR loop. Account locks and a fresh posting check avoid double credit in the same month; the complete batch shares one transaction and has no per-row COMMIT.',
              'sp_batch_mature_fixed_deposits':'Matured status plus the enclosing transaction prevents repeated payouts. The chosen primary savings account is oldest by OpeningDate then AccountNo, restricted to Active. A missing account aborts prior loop effects as well as the current FD.',
            }
            r.para('Integrity and transaction benefit: '+benefits[name])
    r.heading('8.7 SQL Security Implementation',2)
    r.para('The deployed sql/06_security_concurrency.sql hardens the Review 2 logic. It replaces sp_transfer_funds with a SECURITY INVOKER wrapper calling the narrow SECURITY DEFINER _bank_transfer_atomic helper, then COMMIT. The COMMIT-bearing wrapper cannot itself be SECURITY DEFINER or have a SET clause. The helper uses explicit public names and a fixed pg_catalog, public, pg_temp search_path, is owned by the non-superuser bank_admin role, and is explicitly granted only for the transfer operation.')
    r.para('The following source is the deployed wrapper and helper; it supplements, rather than contradicts, the original full procedure source in Section 8.3.')
    for name in ['_bank_transfer_atomic','sp_transfer_funds']:
        row=next(row for row in r.e['routines']['rows'] if row[0]==name)
        r.code(row[4])
    security=r.e['scripts']['06_security_concurrency.sql']
    create='\n'.join(line for line in security.splitlines() if line.startswith(('DROP ROLE IF EXISTS','CREATE ROLE')))
    grants=security[security.index('-- GRANT matrix (existing objects):'):security.index('-- REVOKE demonstration:')]
    # Include the exact admin grants/ownership operations in the role matrix evidence.
    admin_grants=security[security.index('GRANT ALL PRIVILEGES ON DATABASE'):security.index('-- GRANT matrix (existing objects):')]
    def roles(cell):
        r.code(create,cell); r.para('Observed installed role catalogue; role creation was not rerun against shared cluster roles.',parent=cell,size=10); r.output(r.e['roles'],cell)
    def grant_matrix(cell):
        r.code(admin_grants+'\n'+grants,cell); r.output(r.e['privileges'],cell,sample=3)
    def revoke(cell):
        r.code('GRANT DELETE ON public.Account TO bank_teller;\nREVOKE DELETE ON public.Account FROM bank_teller;',cell)
        r.para('Actual effective DELETE privilege: after GRANT = true; after REVOKE = false. Both statements were verified inside a transaction rolled back to preserve the existing grants.',parent=cell,size=10)
    def least(cell):
        r.code(r.e['rbac_demo']['proof_sql'],cell); r.code(r.e['rbac_demo']['output'],cell,result=True)
    def secure(cell):
        alterations=security[security.index('ALTER PROCEDURE public.sp_pay_loan_emi'):security.index('-- Definer code runs as a non-superuser')]
        r.code(alterations,cell)
        r.code("REVOKE CREATE ON SCHEMA public FROM PUBLIC;\nREVOKE ALL PRIVILEGES ON ALL TABLES IN SCHEMA public FROM PUBLIC;\nREVOKE ALL PRIVILEGES ON ALL SEQUENCES IN SCHEMA public FROM PUBLIC;\nREVOKE EXECUTE ON ALL ROUTINES IN SCHEMA public FROM PUBLIC;\nGRANT SELECT ON public.vw_customer_passbook TO bank_teller;\nGRANT EXECUTE ON PROCEDURE public.sp_transfer_funds\n(INT,INT,NUMERIC,VARCHAR,VARCHAR) TO bank_teller;",cell)
        r.para('Installed privilege evidence: teller Account SELECT/INSERT/UPDATE/DELETE=false; passbook SELECT=true; transfer EXECUTE=true. The helper EXECUTE grant is required by the invoker wrapper. No bank_admin membership is granted to teller or manager.',parent=cell,size=10)
    r.field_table([('Create >= 3 users/roles',roles),('GRANT privileges per role matrix',grant_matrix),('REVOKE demonstration',revoke),('Least privilege proof',least),('Secure exposure',secure)])
    r.image('RBAC'); r.image('RBAC-PROOF')
    r.para('Security boundary: bank_admin, branch_manager and bank_teller are NOLOGIN group roles; a trusted session uses SET ROLE for the proof. Employee.Role is descriptive staff data, not an automatic database login or role membership. FastAPI currently connects as postgres and exposes simulated portal identity selection, so application route authorization and branch-limited access are not claimed to be enforced by these grants.')
    r.heading('8.8 Transactions Indexing and Query Optimization',2)
    r.para('PostgreSQL does not accept literal SAVEPOINT or ROLLBACK TO SAVEPOINT commands inside a PL/pgSQL DO body. The exact script therefore demonstrates literal SQL transaction control at top level, and a separate BEGIN/EXCEPTION block as the supported implicit-subtransaction equivalent.')
    start=security.index('BEGIN;\nSAVEPOINT account_demo;'); end=security.index('-- =====================================================================\n-- 4.',start)
    r.code(security[start:end].strip())
    s=r.e['savepoint']
    r.table(['Observation','Actual value'],[['Balance before',s['before']],['Balance inside savepoint',s['inside_savepoint']],['Balance after rollback',s['after_rollback']],['Audit count before / after',f'{s["audit_before"]} / {s["audit_after"]}']],[3.8,3.2])
    r.image('SAVEPOINT')
    r.para('The literal proof ended with COMMIT after restoring the row and its audit. An EXCEPTION block rolls back only its inner changes and permits the outer block to continue. Transactional rollback does not reclaim SERIAL sequence values; gaps do not indicate missing business rows.')
    r.para('Three project indexes',bold=True)
    for name,definition in r.e['indexes']['rows']: r.code(definition+';')
    r.para('The transaction index leads with AccountNo for equality, then TransactionDate for range/ordering. The customer name index matches last-name/first-name access. The loan index supports customer equality followed by status. Indexes consume storage and add maintenance work to writes, so they are not an unconditional improvement.')
    r.code('BEGIN;\nDROP INDEX IF EXISTS public.idx_txn_account_date;\nEXPLAIN (ANALYZE, BUFFERS)\n'+r.e['plan_sql']+'\nROLLBACK;\n-- Original index is restored by transactional rollback.\nEXPLAIN (ANALYZE, BUFFERS)\n'+r.e['plan_sql'])
    r.para('Actual plan without idx_txn_account_date',bold=True)
    r.code('\n'.join(row[0] for row in r.e['plan_without']['rows']),result=True)
    r.para('Actual plan with idx_txn_account_date available',bold=True)
    r.code('\n'.join(row[0] for row in r.e['plan_with']['rows']),result=True)
    r.image('INDEX')
    before='\n'.join(row[0] for row in r.e['plan_without']['rows'])
    after='\n'.join(row[0] for row in r.e['plan_with']['rows'])
    before_time=re.search(r'Execution Time: ([\d.]+ ms)',before)[1]
    after_time=re.search(r'Execution Time: ([\d.]+ ms)',after)[1]
    method='Bitmap Index Scan and Bitmap Heap Scan' if 'Bitmap Index Scan' in after else 'Index Scan' if 'Index Scan' in after else 'a sequential scan'
    r.para(f'The captured plans above use the same predicate and return the same qualifying rows. With the named index available, PostgreSQL chooses {method}; the plan explicitly shows whether a sort remains. Measured execution is {before_time} without the named index versus {after_time} with it available. Compare actual rows, rows removed, heap blocks and shared buffer hits rather than cost estimates alone. These timings are a single cache-sensitive sample, not a statistically valid speedup claim. Reruns and larger representative data are needed for performance conclusions; choosing a sequential scan for a small table can be correct.')
    r.heading('8.9 Review 2 Requirements Checklist',2)
    evidence=[
      ('CREATE TABLE scripts and constraints','sql/01_ddl.sql','Twelve relations and eleven foreign keys; clean disposable rebuild committed; live constraints and definitions in Section 6.'),
      ('INSERT UPDATE DELETE','sql/02_dml.sql; seed_bulk_data.py','DML cycle executed during clean seed; all twelve live counts recorded; atomic bulk loader inspected.'),
      ('Query categories','sql/03_standard_queries.sql','Exactly 35 statements executed, with actual outputs in Section 7; all A-G functions covered, including combined Category F examples.'),
      ('Views','sql/05_advanced_logic.sql','Three live views, catalogue updatability evidence and output samples; four-table passbook, aggregate summary and writable restricted view.'),
      ('COMMIT ROLLBACK SAVEPOINT','sql/05_advanced_logic.sql; sql/06_security_concurrency.sql','Committed isolated transfer/EMI; live before/after SAVEPOINT restoration; implicit subtransaction proof.'),
      ('Procedures and functions','sql/05_advanced_logic.sql','Two core procedures, two business UDFs and two additional cursor procedures installed and exercised.'),
      ('Triggers','sql/05_advanced_logic.sql; sql/07_test_cases.sql','Two triggers and required support functions; TC-08/09 verify old/new audit growth and no-op suppression; repayment changes debt.'),
      ('Indexes and EXPLAIN comparison','sql/06_security_concurrency.sql','Three installed indexes and identical before/after live EXPLAIN ANALYZE BUFFERS plans reproduced.'),
      ('Execution evidence','sql/03_standard_queries.sql; sql/04_complex_queries.sql; sql/07_test_cases.sql','Live SQL outputs and exact PostgreSQL NOTICE text embedded; matching DBeaver screenshots accompany the shared SQL demonstrations.' if r.screenshots else 'Live SQL outputs and exact PostgreSQL NOTICE text embedded; labelled screenshot slots allow optional DBeaver exhibits.'),
      ('Design changes since Review 1','sql/01_ddl.sql; PROJECT_SPEC.md','Section 6.1 change log with dependency reasons; generated-column formal qualification is explicit.'),
    ]
    r.table(['#','Mandatory component','Script and concrete evidence','Status'],
            [[str(i),name,script+'\n'+detail,'Completed & Verified'] for i,(name,script,detail) in enumerate(evidence,1)], [.35,1.5,3.9,1.25],font=9)
    r.para(('Checklist status refers to implementation and embedded execution evidence. '+('Review 3 includes genuine DBeaver and React screenshots captured from the running local project. ' if r.review==3 else 'Review 2 includes genuine DBeaver screenshots of the shared database implementation. ')+'Longer statements and result sets are partial viewport previews; the full SQL and sampled output are printed alongside them. The Chen ER model and complete attribute panels are embedded in Section 5. ' if r.screenshots else 'Checklist status refers to implementation and embedded execution evidence. Manual UI/DBeaver screenshot exhibits are explicitly marked placeholders and are not represented as captured images. ')+'All A-G query counts meet their required ranges, including four Category F queries within the permitted four-to-six range.')


def work_distribution(r):
    r.heading('9 Work break down',newpage=True)
    r.para('The responsibility allocation is approximately 60% / 20% / 20%. Architecture and transactional correctness are concentrated in the first role; query coverage, datasets, interface modules and verification provide separate, reviewable contributions for the other two roles. Percentages describe assigned scope, not measured hours.')
    assignments2=[
       '60% Core architecture and heavy engineering. Twelve-table schema redesign and normalization; generated financial columns; sp_transfer_funds and sp_pay_loan_emi; trg_audit_account_changes and trg_loan_repayment; both explicit cursor batches; SAVEPOINT/ROLLBACK transaction design; three indexes and EXPLAIN ANALYZE comparison.',
       '20% Querying and data foundations. Categories A-D: basic, aggregate WHERE/HAVING contrast, all join forms and subqueries including the NOT IN NULL pitfall and fix. Three database views and atomic Faker bulk seeding through seed_bulk_data.py.',
       '20% Analytics and function verification. Categories E-G: set operations, advanced built-ins and window functions. CQ-01 through CQ-10 business-query verification; fn_calculate_emi and fn_customer_net_worth implementation and reference checks.',
    ]
    additions=[
       ' Review 3 extension: all Review 2 architecture plus two-session FOR UPDATE concurrency, database RBAC architecture, FastAPI transaction engine and core React transactional API integration.',
       ' Review 3 extension: Customer NetBanking passbook, transfers, beneficiaries and FD Vault screens; execute and document TC-01 through TC-08.',
       ' Review 3 extension: official executive analytics, customer 360, loan origination and audit screens; execute and document TC-09 through TC-15.',
    ]
    r.table(['Registration Number','Name','Work Assigned'],
            [[reg,name,work+(additions[i] if r.review==3 else '')] for i,((name,reg,phone),work) in enumerate(zip(MEMBERS,assignments2))], [1.2,1.8,4.0],font=10)
    r.para('Continuity: Review 3 inherits the same schema, query catalogue and database objects documented for Review 2. Final application and concurrency work extend those responsibilities rather than replacing or reassigning the earlier database foundation.')


def review3(r):
    r.heading('10 Review Requirements',newpage=True)
    r.para('The final review demonstrates an end-to-end banking workflow backed by the same twelve-table design. The verification evidence below separates installed database security, transaction and trigger correctness, query plans, concurrency and the visible React modules. The operational dataset is retained while reproducible destructive setup is demonstrated only in a dedicated evaluation database.')
    r.heading('10.1 Concurrency and Locking Demonstration',2)
    r.para('Two real PostgreSQL connections were used. Session A locked account 100001 without changing it. Session B attempted FOR UPDATE NOWAIT on the same row and received SQLSTATE 55P03. Both sessions then rolled back, releasing all locks. Without NOWAIT, the second writer waits and, at Read Committed, reads the current committed row after acquiring its lock.')
    r.code('-- Session A\nBEGIN;\nSELECT AccountNo, Balance FROM public.Account\nWHERE AccountNo=100001 FOR UPDATE;\n-- Keep transaction open while Session B runs.\nROLLBACK;\n\n-- Session B while A holds the row\nDO $$ BEGIN\n  PERFORM AccountNo FROM public.Account\n  WHERE AccountNo=100001 FOR UPDATE NOWAIT;\n  RAISE EXCEPTION \'Expected a conflicting row lock\';\nEXCEPTION WHEN lock_not_available THEN\n  RAISE NOTICE \'PASS: second writer blocked by first writer (55P03)\';\nEND $$;')
    c=r.e['concurrency']
    r.table(['Session','Actual observation'],[['A',f'Locked AccountNo={c["session_a"][0]}, Balance={c["session_a"][1]}'],['B',c['session_b_error']+'; SQLSTATE '+c['sqlstate']],['Result','PASS; both sessions rolled back, no account row changed.']],[1,6])
    r.image('CONCURRENCY')
    r.para('Lost-update reasoning: if two writers read balance B before locking, then each writes B minus its own payment, one debit can overwrite the other. Our cash and procedure workflows acquire row locks before reading the balance, then validate and post inside the same transaction. Internal transfers lock both account numbers in ascending order; EMI locks Account before Loan. Atomic UPDATE Balance=Balance-amount also avoids a stale application-supplied absolute value, but must still be coupled with an atomic sufficient-funds check.')
    r.heading('10.2 Comprehensive Verification Test Suite',2)
    r.para('Every exact TC block from sql/07_test_cases.sql was executed against banking_db. Intentional errors are caught by nested EXCEPTION handlers that check the expected SQLSTATE and, where relevant, the exact constraint. A missing rejection or unrelated failure raises a fatal test error. Successful audit-test changes are deliberately rolled back. All fifteen cases passed; sequence numbers may advance despite row rollback.')
    actions=[
      ('Primary key','INSERT a copy of CustomerID 1','23505 customer_pkey rejects duplicate'),
      ('Foreign key','Set Account 100001 CustomerID=-2147483648','23503 account_customerid_fkey'),
      ('Foreign key','Set Loan 2001 BranchID=-2147483648','23503 loan_branchid_fkey'),
      ('CHECK age','Set Customer 1 Age=17','23514 customer_age_check'),
      ('CHECK balance','Set Account 100001 Balance=-1','23514 account_balance_check'),
      ('CHECK gender',"Set Customer 1 Gender='Invalid'",'23514 customer_gender_check'),
      ('UNIQUE PAN','Give Customer 2 Customer 1 PAN','23505 customer_pan_number_key'),
      ('Audit credit','Credit 1 to Account 100001; check exact old/new text, then undo','One new correct audit inside test; no retained change'),
      ('Audit debit and no-op','Debit 0.01 from 100003 then self-assign balance; undo','One audit for real debit; none for no-op; rollback restores both'),
      ('Procedure rejection','Transfer sender balance+1 to 100003','22023 insufficient balance; balances/history/audits unchanged'),
      ('Passbook view','Compare view row count with four-table join','Nonempty matching count; no displayed NULL Remarks'),
      ('Summary view','Reconcile branch counts, loans and attributed FDs','One row per branch; no independent-product fan-out'),
      ('EMI UDF','Reference EMI, zero rate and every stored loan','21695.58 reference; 1000 zero-rate; generated values agree'),
      ('Net worth UDF','Compare every customer to signed UNION ALL products','All customers reconcile, including missing-product NULLs'),
      ('Data type',"Dynamic SQL assigns 'not-an-integer' to Age",'22P02 invalid integer input'),
    ]
    r.table(['Test ID','Test Category','Input / Action','Expected Result','Actual PostgreSQL Output','Status'],
            [[t['id'],*action,t['output'],'PASS'] for t,action in zip(r.e['tests'],actions)], [.5,.8,1.5,1.45,2.25,.5],font=8.5)
    for test in r.e['tests']:
        r.image(test['id'])
    r.para('The original executable test file remains the authoritative input. For an evaluator who wants to see the wrapped failure pattern, TC-05 is reproduced below; the remaining blocks follow the same guarded structure with test-specific assertions.')
    r.code(next(t['sql'] for t in r.e['tests'] if t['id']=='TC-05'))
    r.heading('10.3 Live Demonstration Runbook',2)
    r.para('Allocate approximately eleven minutes. Use a dedicated disposable evaluation database for the clean DDL/DML step: sql/01_ddl.sql drops data and dependants. Do not point the reset at the populated banking_db shown in the application. Keep two SQL sessions ready and an API/React terminal running before the viva.')
    r.table(['Step and time','Presenter action','Concrete evidence to show'],[
      ['1 Clean DDL and DML\n0:00-1:30','Create a separate evaluation database; run 01_ddl.sql and 02_dml.sql with ON_ERROR_STOP; install 05 definitions/verification before 06. The report evidence collector already proves the clean 01/02/05 sequence.','Twelve relations, constraints, DML cycle and seed counts. Use the live row-count table to distinguish the populated operational snapshot.'],
      ['2 Role based access\n1:30-3:00','Use the installed roles on banking_db. SET ROLE bank_teller; SELECT * FROM vw_customer_passbook LIMIT 3; attempt DELETE FROM Account WHERE false in the guarded proof; RESET ROLE.','Passbook allowed, direct DELETE denied with 42501. Show the effective privilege matrix. Do not rerun the cluster-wide role reset in a second database.'],
      ['3 Transaction and deliberate failure\n3:00-5:00','Choose an owned Active sender and receiver. Demonstrate a valid transfer on disposable fixtures; then request sender balance+1 or use a Frozen receiver. Run TC-10 for a repeatable wrapped rejection.','Commit notification for the valid workflow; exact SQL exception for rejection; both balances/history/audit unchanged on failure.'],
      ['4 Trigger firing\n5:00-6:30','Run TC-08/09; show audit growth during the update and restoration after rollback. Demonstrate EMI in the verification dataset and inspect Loan_Repayment plus Loan outstanding.','One old/new balance audit; no-op creates none. A Success repayment changes debt, Pending/Failed does not.'],
      ['5 Index comparison\n6:30-8:30','Run the identical EXPLAIN (ANALYZE, BUFFERS) predicate first with the named index dropped inside BEGIN/ROLLBACK, then with it restored.','Plan node change, actual row counts, buffer hits and timings. Explain small-table/cache limits; do not force a speedup.'],
      ['6 Two session concurrency\n8:30-11:00','Session A locks Account 100001 FOR UPDATE; B tries FOR UPDATE NOWAIT. Show 55P03; rollback A and retry B.','Second writer cannot act on the locked row; explain the stale-read lost-update example and deterministic lock order. End with both transactions closed.'],
    ],[1.15,3.45,2.4],font=9.5)
    r.code('# Example safe clean setup in PowerShell\ndocker exec banking-postgres createdb -U postgres banking_viva\nGet-Content -Raw sql/01_ddl.sql | docker exec -i banking-postgres psql -X -v ON_ERROR_STOP=1 -U postgres -d banking_viva\nGet-Content -Raw sql/02_dml.sql | docker exec -i banking-postgres psql -X -v ON_ERROR_STOP=1 -U postgres -d banking_viva\nGet-Content -Raw sql/05_advanced_logic.sql | docker exec -i banking-postgres psql -X -v ON_ERROR_STOP=1 -U postgres -d banking_viva\n# Shared role grants are already installed on banking_db by 06.\n# Use their guarded proof there, rather than resetting cluster roles here.')
    r.heading('10.4 Working Application Modules',2)
    r.para('The frontend is React 19 with Vite, Tailwind CSS v4, Recharts and Lucide; Axios calls FastAPI on port 8000. Customer and official dashboards have six and seven distinct tabs. A shared sun/moon switch follows the system preference initially, persists explicit choice in localStorage and themes charts, tables, inputs, dialogs and notifications. Navigation can switch real customer or employee profiles instantly.')
    r.heading('Smart Login and Global Identity Switcher',3)
    r.para('Inputs: Customer NetBanking or Bank Official role; searchable real customer directory showing IDs, owned accounts, balances and KYC; direct CustomerID/AccountNo option; employee or system-admin choice. API calls: GET /api/directory/customers-preview, /employees-preview and /lookups, followed by POST /api/login. Output: selected profile and portal route, saved locally. The switcher uses the same real directory without guessing IDs.')
    r.placeholder('[PLACEHOLDER: Insert UI Screenshot - Smart Login and Identity Switcher]')
    for portal,modules in [('Customer NetBanking Portal',CUSTOMER_MODULES),('Bank Official Portal',ADMIN_MODULES)]:
        r.heading(portal,3)
        for i,(name,inputs,calls,outputs) in enumerate(modules,1):
            r.para(f'Tab {i}  {name}',bold=True)
            r.para('Accessible inputs: '+inputs)
            r.para('Backend API and SQL: '+calls)
            r.para('Visible outputs: '+outputs)
            r.placeholder(f'[PLACEHOLDER: Insert UI Screenshot - {name}]')
            if name=='My Accounts and Passbook' and 'Passbook Statement Detail' in r.ui_screenshots:
                r.image_record(r.ui_screenshots['Passbook Statement Detail'],'Passbook statement rows')
            if name=='Batch Cursors and Live Trigger Audit' and 'Live Audit Trail Detail' in r.ui_screenshots:
                r.image_record(r.ui_screenshots['Live Audit Trail Detail'],'Live audit rows')
    r.para('Transaction contract: request amounts are validated as finite decimal values, sent through Psycopg2 parameters and returned as decimal strings. Direct sp_transfer_funds CALL uses its own autocommit connection for explicit COMMIT; multi-statement application workflows use a single managed transaction. PostgreSQL failures retain their original text and SQLSTATE for the review demonstration. Validation failures are 422, ownership mismatches can return 403, and absent records return 404.')
    r.para('API security qualification: current portal route guards and saved identities are navigation conveniences, not authenticated authorization. No passwords are stored by the app. Development CORS is unrestricted, and the API uses the supplied local postgres account. Production deployment would require authenticated identities, role/branch authorization, a least-privileged service account, secret management and an intentional error-disclosure policy; those are not represented as delivered security controls.')
    r.heading('10.5 Viva Preparation and Evaluation Anchors',2)
    viva=[
      ('Why use AccountNo as the primary key?','It identifies one account independently of its owner and branch. CustomerID cannot be the account key because one customer can hold multiple accounts.'),
      ('Show a functional dependency and closure.','In Postal_Directory, Pincode -> City,State, so Pincode+ is {Pincode,City,State}; it is a candidate key under the canonical-postal-location assumption.'),
      ('Which update anomaly was removed?','Repeating City/State in Branch would allow two branches with one pincode to disagree after an update. The lookup holds the canonical fact once.'),
      ('Are the generated-column tables strictly 3NF or BCNF?','Their independent base facts are normalized; their stored generated extensions have non-key determinants. Section 6 states this qualification rather than asserting BCNF.'),
      ('When might you denormalize?','A cached branch aggregate may help high-volume reporting, but introduces refresh/consistency costs. Our view aggregates product totals without a separate persistent summary table.'),
      ('Where is the M:N junction?','This model has no joint-account or arbitrary customer-product junction. Customer-to-Branch can be traversed through Account, whose AccountNo distinguishes multiple accounts for one owner/branch pair.'),
      ('Why only selected delete cascades?','Customer -> Account/Beneficiary and Loan -> Loan_Repayment are declared cascades. Transaction history can still block account deletion. Broad cascades could destroy financial evidence.'),
      ('What if a referenced parent is deleted?','The FK NO ACTION rule rejects deletion while protected child rows exist, unless the FK explicitly declares CASCADE. Demonstrate only in a rollback transaction.'),
      ('Why combine CHECK and NOT NULL?','CHECK rejects false but accepts unknown. Balance is NOT NULL plus CHECK Balance>=0; a nullable status CHECK does not by itself reject NULL.'),
      ('DELETE TRUNCATE or DROP?','DELETE removes qualifying rows and fires row DELETE triggers; TRUNCATE removes all rows with stronger locking and separate trigger rules; DROP removes objects. PostgreSQL can roll these DDL/data effects back, subject to dependency rules.'),
      ('Why LEFT JOIN in a business report?','CQ-09 retains loans without any repayment attempt; an INNER JOIN would omit them and hide zero-collection loans.'),
      ('WHERE versus HAVING?','B-04 applies Active status before grouping, then tests the branch SUM in HAVING. SUM cannot be used in that WHERE level.'),
      ('Why does NOT IN fail with NULL?','An unmatched comparison against a set containing NULL is UNKNOWN. D-06 filters the inner NULLs; NOT EXISTS is another intentional anti-join design.'),
      ('EXISTS versus IN?','EXISTS tests row existence without multiplying the outer customer; D-04 uses it with correlated criteria. PostgreSQL may transform either form into a semi-join; measure the actual plan.'),
      ('Correlated subquery or join on large data?','A join rewrite must preserve multiplicity, NULL and aggregation semantics. Neither syntax is always faster; compare representative plans, estimates and buffers.'),
      ('UNION or UNION ALL?','UNION removes duplicates; UNION ALL retains individual product rows in E-02. Removing duplicates can destroy intended multiplicity and adds distinct-processing work.'),
      ('COUNT star column and DISTINCT?','COUNT(*) counts rows, COUNT(Remarks) skips NULL remarks, COUNT(DISTINCT Remarks) counts distinct non-NULL values. A-06 proves missing remarks exist.'),
      ('Why index AccountNo before TransactionDate?','The demonstration uses equality on account and a date range. A predicate only on the trailing date has different selectivity/access costs and is not guaranteed to use this index.'),
      ('When can an index be ignored?','For a small table or a predicate matching most rows, sequential scanning can cost less than index access plus heap fetches. Planner choice is not an implementation failure.'),
      ('Explain the measured plan.','The baseline scans 642 rows and retains 21; the indexed plan uses a bitmap index then heap scan and sorts by date. Actual timings are too small and cache-sensitive for a general speedup claim.'),
      ('How to prove an index benefit?','Use larger representative datasets, maintained statistics, repeated measurements and buffer/row analysis with identical predicates. Do not force enable_seqscan off as proof of a natural planner benefit.'),
      ('What do more indexes cost?','Every INSERT/DELETE and indexed-column UPDATE maintains index entries; storage, WAL, vacuum and cache pressure can grow. Add indexes for measured access patterns.'),
      ('What prevents two debits from overspending?','Lock the account before reading and checking its balance, then debit in the same transaction. A competing writer waits; an unlocked stale application read is insufficient.'),
      ('Two ways to prevent a lost update?','Use SELECT FOR UPDATE before read-modify-write, or a conditional atomic UPDATE with a Balance>=amount predicate and checked row count. Serializable transactions can also require application retries.'),
      ('What anomalies remain at Read Committed?','PostgreSQL never exposes dirty reads. Read Committed uses a new snapshot per statement, so non-repeatable reads and phantoms can occur. Repeatable Read removes those snapshot changes but can still need serialization reasoning.'),
      ('How is committed data durable?','WAL and commit flush behavior, checkpoints, fsync and synchronous_commit settings determine recovery guarantees. The project relies on PostgreSQL defaults and does not claim a separately tested power-loss recovery experiment.'),
      ('How can deadlocks occur?','Opposite transfers locking sender first can form a cycle. Our transfer locks both Account numbers in one ascending order; clients must still handle retryable deadlock/serialization errors.'),
      ('Procedure function or trigger?','EMI calculation is a SELECT-callable function; transfer is a CALL procedure with transaction control. Auditing and repayment side effects are triggers. Hidden trigger recursion or unbounded side effects would be hazards.'),
      ('Why keep explicit cursors?','A set-based UPDATE is often more efficient, but the rubric requires explicit iteration. Our bound cursor loops illustrate per-row checks and coordinated audit/history postings within one atomic batch.'),
      ('How are injection branch access and passwords handled?','Psycopg2 parameters separate values from SQL; RBAC does not replace parameterization. The current API has simulated identity and no password store or enforced branch authorization. A production design would authenticate users and use branch-filtered policies with least-privileged connections.'),
      ('How would you recover an accidental deletion?','A verified base backup plus archived WAL can support point-in-time recovery before the deletion. A routine pg_dump alone is not arbitrary PITR. No archived-WAL/PITR environment is implemented or tested in this lab.'),
    ]
    for i,(q,a) in enumerate(viva,1):
        r.para(f'{i}. {q}',bold=True); r.para(a)


def final_evaluation(r):
    number=11 if r.review==3 else 10
    r.heading(f'{number} REVIEW EVALUATION',newpage=True)
    r.para('Faculty assessment. The original grading grid is preserved below; individual scores and signatures are reserved for the evaluator.',size=10)
    sect=r.d.element.body.find(qn('w:sectPr'))
    sect.addprevious(deepcopy(r.grading))
    r.fill_contents()
    # Basic metadata is intentional; the original header, footer and media parts remain.
    r.d.core_properties.title=TITLE
    r.d.core_properties.subject=f'DATABASE SYSTEMS BCSE302P Review {r.review}'
    r.d.core_properties.author='; '.join(x[0] for x in MEMBERS)
    r.d.core_properties.keywords='PostgreSQL; Core Banking; BCSE302P; Review '+str(r.review)


def ensure_backups():
    folder=ROOT/'report_backups'/'original'; folder.mkdir(parents=True,exist_ok=True)
    for review in [2,3]:
        name=f'Project Template - review {review}.docx'
        target=folder/name
        if not target.exists(): shutil.copy2(ROOT/name,target)
    return folder


def write_contract(backups):
    inventory_file=WORK/'template_inventory.json'
    if inventory_file.exists():
        inventory=json.loads(inventory_file.read_text(encoding='utf-8'))
    else:
        inventory={}
        for source in backups.glob('*.docx'):
            with ZipFile(source) as package:
                parts=[{'path':name,'size':len(package.read(name)), 'sha256':hashlib.sha256(package.read(name)).hexdigest()} for name in package.namelist()]
            inventory[source.name]={'sha256':hashlib.sha256(source.read_bytes()).hexdigest(),
                                    'images':[part for part in parts if 'media/' in part['path']], 'inventory':parts}
        inventory_file.write_text(json.dumps(inventory,indent=2),encoding='utf-8')
    text=['# Review report template contract','',
          'Reference templates are retained unchanged in report_backups/original. Both reports use Letter portrait pages, 0.75 inch top/bottom, 0.8 inch left and 0.7 inch right margins; one section per document. The inherited header and PAGE footer remain unchanged.',
          'Cover paragraphs 0-20 are retained, including the VIT drawing at paragraph 20. Editable slots: project title, review label, three names/registrations, course and recipient. The full project title uses an 18 pt Title paragraph at the source centered position; content otherwise retains the centered source layout.',
          'Body slots after the cover are replaced with the requested banking report. Source Times New Roman 11 pt body and 14/12 pt heading hierarchy are reused; all headings are black. The obsolete sample ER/model images remain package-preserved but are not displayed. Section 5 now contains the user-authorized browser-rendered Mermaid Chen overview and four complete attribute panels.',
          'Body content may expand to any number of pages. Source section order is retained through Section 8; Section 9 holds the work allocation; Review 3 Section 10 holds final requirements. Grading tables are copied exactly and placed last. Empty grading cells remain evaluator-owned.',
          'The old stale TOC content control is replaced by a stable source-derived text TOC with tab leaders and verified final page numbers. Header/footer, numbering relationships, and all embedded image binaries are preserve-only. New body tables use visible light-gray borders, blue-gray headers, repeating header rows, natural row height and explicit widths. SQL uses Consolas 9 pt shaded paragraphs.',
          'Reference renders are .report_work/template-review2.pdf and template-review3.pdf with all-page PNGs. Packaged LibreOffice renderer failed because no soffice exists; installed Word 2024 provides the read-only PDF render and bundled Poppler rasterizes pages. Final fidelity gates include logo hashes, exact grading-table XML, page geometry, every-page rendered QA, labels, source counts and live-evidence provenance.','']
    for name,info in inventory.items():
        digest=hashlib.sha256((backups/name).read_bytes()).hexdigest()
        if digest!=info['sha256']: raise RuntimeError('Retained source hash mismatch: '+name)
        text += [f'## {name}',f'Reference: {backups/name}',f'SHA256: {digest}',f'Original pages: {9 if "2.docx" in name else 10}; sections: 1.',
                 'Preserve-only binary image inventory: '+json.dumps(info['images']),
                 'Full style, relationship and part inventory: template_inventory.json.','']
    (WORK/'artifact.md').write_text('\n'.join(text),encoding='utf-8')


def verify_package(reference,output):
    with ZipFile(reference) as a,ZipFile(output) as b:
        for name in a.namelist():
            if name.startswith(('word/media/','word/header','word/footer')):
                if a.read(name)!=b.read(name): raise RuntimeError('Preserved package part changed: '+name)
    src=Document(reference); dst=Document(output)
    if src.tables[-1]._tbl.xml!=dst.tables[-1]._tbl.xml: raise RuntimeError('Grading grid was modified')
    if len(dst.sections)!=1: raise RuntimeError('Unexpected section count')
    for attr in ['page_width','page_height','top_margin','bottom_margin','left_margin','right_margin']:
        if getattr(src.sections[0],attr)!=getattr(dst.sections[0],attr): raise RuntimeError('Page geometry changed')
    text='\n'.join(p.text for p in dst.paragraphs)
    if 'Malpractice Management' in text or 'Group Member1' in text: raise RuntimeError('Template sample text remains')
    return len(dst.paragraphs),len(dst.tables)


def preserve_opaque_parts(reference,output):
    """python-docx may reserialize untouched header XML; restore it verbatim."""
    temporary=output.with_suffix('.preserved.docx')
    with ZipFile(reference) as original, ZipFile(output) as generated, ZipFile(temporary,'w') as final:
        for info in generated.infolist():
            name=info.filename
            preserve=name.startswith(('word/media/','word/header','word/footer','word/_rels/header','word/_rels/footer'))
            final.writestr(info,original.read(name) if preserve and name in original.namelist() else generated.read(name))
    temporary.replace(output)


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--refresh-evidence',action='store_true')
    parser.add_argument('--page-map',type=Path)
    parser.add_argument('--review',type=int,choices=[2,3],help='Update only the requested review; the other file stays untouched.')
    parser.add_argument('--evidence',type=Path,default=WORK/'evidence.json')
    parser.add_argument('--output-dir',type=Path,help='Write a review copy here when the original is open in Word.')
    args=parser.parse_args()
    if args.refresh_evidence or not (WORK/'evidence.json').exists():
        subprocess.run(['uv','run','python',str(ROOT/'collect_report_evidence.py')],cwd=ROOT,check=True)
    evidence=json.loads(args.evidence.read_text(encoding='utf-8'))
    backups=ensure_backups(); write_contract(backups)
    map_file=args.page_map or WORK/'page_map.json'
    page_map=json.loads(map_file.read_text(encoding='utf-8')) if map_file.exists() else {}
    for review in ([args.review] if args.review else [2,3]):
        name=f'Project Template - review {review}.docx'
        r=Report(backups/name,evidence,review,page_map)
        abstract_and_scope(r); schema_section(r); standard_section(r); advanced_section(r); work_distribution(r)
        if review==3: review3(r)
        final_evaluation(r)
        output_dir=args.output_dir or ROOT
        output_dir.mkdir(parents=True,exist_ok=True)
        output=output_dir/name
        if output.exists() and (ROOT/'report_backups'/'latest').exists():
            shutil.copy2(output,ROOT/'report_backups'/'latest'/name)
        r.d.save(output)
        preserve_opaque_parts(backups/name,output)
        counts=verify_package(backups/name,output)
        print(f'Review {review} populated: {output} ({counts[0]} paragraphs, {counts[1]} tables); cover media and grading grid preserved.')
    print('Capture final pagination with render_reports.ps1, then regenerate with --page-map .report_work/page_map.json and render again.')


if __name__=='__main__': main()
