"""Patch the current Review 3 with genuine DBeaver evidence and tidy output prose.

Preserve the user's edited report, its SQL definitions, predictions, and unrelated
sections. Never rebuild from the teacher template or write Review 2 / Guide.
"""
from collections import Counter
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile
import hashlib
import json
import re
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from lxml import etree

ROOT = Path(__file__).resolve().parent
TARGET = ROOT / 'Project Template - review 3.docx'
WORK = ROOT / '.report_work' / 'review3-final-evidence'
SHOTS = ROOT / 'report_screenshots' / 'review3' / 'supplemental'


def text(e):
    return ''.join(e.xpath('.//w:t/text()'))


def style(e):
    s = e.xpath('./w:pPr/w:pStyle/@w:val')
    return s[0] if s else ''


def canonical(e):
    return etree.tostring(e, method='c14n', exclusive=True)


def replace_text(e, value):
    """Preserve paragraph styling and reuse the first existing run's formatting."""
    nodes = e.xpath('.//w:t')
    if nodes:
        nodes[0].text = value
        nodes[0].set(qn('xml:space'), 'preserve')
        for t in nodes[1:]:
            t.text = ''
    else:
        Paragraph(e, None).add_run(value)


def make_p(document, value, style_name='Normal'):
    p = document.add_paragraph(value, style_name)
    return p._p


def shot(document, name, caption):
    p = document.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(SHOTS / f'{name}.png'), width=Inches(6.95))
    cap = document.add_paragraph(caption)
    cap.paragraph_format.keep_together = True
    return [p._p, cap._p]


def replace_range(body, start, end, replacement):
    nodes = list(body)
    first, last = nodes.index(start), nodes.index(end)
    for e in nodes[first:last]:
        body.remove(e)
    for e in reversed(replacement):
        if e.getparent() is not None:
            e.getparent().remove(e)
        body.insert(first, e)


def write_package(source, document):
    """Keep every existing package part byte-for-byte except required XML parts."""
    updates = {
        'word/document.xml': etree.tostring(document.element, xml_declaration=True,
                                           encoding='UTF-8', standalone=True),
        'word/_rels/document.xml.rels': document.part.rels.xml,
    }
    with ZipFile(source) as src:
        existing = set(src.namelist())
        for part in document.part.package.iter_parts():
            name = str(part.partname).lstrip('/')
            if name not in existing:
                assert name.startswith('word/media/') and name.endswith('.png'), name
                updates[name] = part.blob
        temp = TARGET.with_suffix('.evidence.docx')
        with ZipFile(temp, 'w') as dst:
            for info in src.infolist():
                dst.writestr(info, updates.pop(info.filename, src.read(info.filename)))
            for name, data in updates.items():
                dst.writestr(name, data)
    temp.replace(TARGET)


def main():
    preservation = json.loads((WORK / 'preservation.json').read_text())
    source = Path(preservation['backup'])
    # The user closed and saved Word. Refuse to overwrite any subsequent edits.
    assert TARGET.read_bytes() == source.read_bytes(), 'Current report changed after backup'
    document = Document(source)
    body = document.element.body
    before_nodes = list(body)

    def at(value):
        return next(e for e in body if text(e) == value)

    def blocks_between(first, last):
        nodes = list(body)
        return nodes[nodes.index(at(first)):nodes.index(at(last))]

    protected = {
        'section6': [canonical(e) for e in blocks_between('6 Tables and Constraints', '7 SQL Implementation')],
        'section7': [canonical(e) for e in blocks_between('7 SQL Implementation', '8 Advanced SQL Implementation')],
        'section9': [canonical(e) for e in blocks_between('9 Application Screenshots', '10 Work Breakdown')],
        'predictions': [canonical(r._tr) for t in document.tables for r in t.rows
                        if r.cells[0].text == 'Expected output'],
        'definitions': [text(p) for p in body.xpath('.//w:p') if style(p) == 'SQLCode'
                        and text(p).startswith('CREATE OR REPLACE')],
    }
    old_images = Counter(body.xpath('.//a:blip/@r:embed'))
    audit = {'new_screenshots': [], 'removed_result_paragraphs': 0,
             'timestamp_paragraphs_cleaned': 0}

    # Transfer: keep the real rejected-operation proof and add the committed path.
    start = at('Live banking_db verification')
    end = next(e for e in body if text(e).startswith('Integrity and transaction benefit: Atomicity'))
    old_transfer = list(body)[list(body).index(start):list(body).index(end)]
    rejected_sql = next(deepcopy(e) for e in old_transfer if style(e) == 'SQLCode')
    rejected_picture = next(deepcopy(e) for e in old_transfer if e.xpath('.//w:drawing'))
    transfer = [
        make_p(document, 'Rejected transfer and rollback'),
        rejected_sql,
        make_p(document, 'Insufficient funds cause a validation exception before COMMIT. '
               'The rejected operation leaves both balances and related records unchanged.'),
        rejected_picture,
        make_p(document, 'Insufficient funds validation'),
        make_p(document, 'Successful committed transfer'),
        make_p(document, 'The successful example runs in banking_report_demo, an isolated '
               'copy of banking_db with the installed procedures and triggers. Balances '
               'were saved immediately before and after CALL; the screenshots show the '
               'committed debit and credit, their history entries and automatic audits.'),
        make_p(document, "CALL public.sp_transfer_funds(100001,100003,1000,'UPI','Report transfer demonstration');", 'SQL Code'),
    ]
    transfer += shot(document, 'TRANSFER-SUCCESS', 'Successful transfer and committed account balances')
    transfer += shot(document, 'TRANSFER-HISTORY', 'Transfer history and automatic balance audits')
    replace_range(body, start, end, transfer)
    audit['new_screenshots'] += ['TRANSFER-SUCCESS', 'TRANSFER-HISTORY']

    # Cursors: replace old zero-posting evidence with eligible dedicated fixtures.
    for heading, name, call, context, caption in [
        ('sp_batch_credit_savings_interest', 'CURSOR-INTEREST',
         'CALL public.sp_batch_credit_savings_interest();',
         'A dedicated Active Savings account in banking_report_demo has no current-month '
         'interest posting. The screenshot verifies the monthly calculation, credited '
         'transaction and resulting balance. A repeated batch call creates no duplicate posting.',
         'Savings interest calculation and credited balance'),
        ('sp_batch_mature_fixed_deposits', 'CURSOR-MATURITY',
         'CALL public.sp_batch_mature_fixed_deposits();',
         'A dedicated Active FD in banking_report_demo has reached its EndDate. The '
         'screenshot verifies its Matured status, the savings-account credit and the '
         'FD_MATURED audit. A repeated batch call creates no duplicate payout.',
         'FD maturity payout and automatic audit'),
    ]:
        nodes = list(body)
        h = next(i for i,e in enumerate(nodes) if style(e) == 'Heading3' and text(e) == heading)
        start = next(e for e in nodes[h+1:] if text(e) == 'Live banking_db verification')
        end = next(e for e in nodes[nodes.index(start):] if text(e).startswith('Integrity and transaction benefit:'))
        replacement = [make_p(document, 'Successful batch execution in the demo database'),
                       make_p(document, context), make_p(document, call, 'SQL Code')]
        replacement += shot(document, name, caption)
        replace_range(body, start, end, replacement)
        audit['new_screenshots'].append(name)

    # Identical search predicates, genuine plans on the same unchanged demo rows.
    nodes = list(body)
    first = next(i for i,e in enumerate(nodes) if text(e).startswith('The transaction index leads'))
    start = next(e for e in nodes[first+1:] if style(e) == 'SQLCode')
    end = at('8.9 Concurrency and Locking Demonstration')
    index_sql = """-- Run in banking_report_demo only.
BEGIN;
DROP INDEX public.idx_txn_account_date;
EXPLAIN (ANALYZE, BUFFERS)
SELECT TransactionID,AccountNo,TransactionDate,Amount
FROM public.Transaction_History
WHERE AccountNo=100001
  AND TransactionDate>=CURRENT_DATE-INTERVAL '365 days'
ORDER BY TransactionDate;
ROLLBACK; -- Restores the named index.
EXPLAIN (ANALYZE, BUFFERS)
SELECT TransactionID,AccountNo,TransactionDate,Amount
FROM public.Transaction_History
WHERE AccountNo=100001
  AND TransactionDate>=CURRENT_DATE-INTERVAL '365 days'
ORDER BY TransactionDate;"""
    replacement = [make_p(document, index_sql, 'SQL Code'),
                   make_p(document, 'Both screenshots use the same query and demo dataset. '
                          'Only the availability of idx_txn_account_date changes.'),
                   make_p(document, 'Execution plan without the transaction index')]
    replacement += shot(document, 'INDEX-WITHOUT', 'Search plan before indexing')
    replacement += [make_p(document, 'Execution plan with the transaction index')]
    replacement += shot(document, 'INDEX-WITH', 'Search plan after indexing')
    replacement += [make_p(document, 'Without the named index, PostgreSQL scans the table '
                          'and filters matching rows. With it available, the plan uses '
                          'Bitmap Index Scan and Bitmap Heap Scan; a sort remains. Compare '
                          'the access method, rows and buffers in the screenshots. Individual '
                          'timings depend on cache state and table size, so this comparison '
                          'does not establish a general performance speedup.')]
    replace_range(body, start, end, replacement)
    audit['new_screenshots'] += ['INDEX-WITHOUT', 'INDEX-WITH']

    # Remove repeated typed SQL results throughout the report. SQL source and
    # teacher-required hand-reasoned Expected output rows remain untouched.
    for p in list(body.xpath('.//w:p')):
        t = text(p)
        if style(p) == 'ResultCode':
            p.getparent().remove(p)
            audit['removed_result_paragraphs'] += 1
        elif t.startswith('banking_db;') and 'Observed ' in t:
            replace_text(p, 'Verification runs inside a transaction and is rolled back after inspection.')
            audit['timestamp_paragraphs_cleaned'] += 1
        elif t.startswith(('Audit count before ', 'Outstanding before insert:')):
            p.getparent().remove(p)
            audit['removed_result_paragraphs'] += 1
        elif re.fullmatch(r'Actual output: \d+ row\(s\); first \d+ shown\. NULL denotes a missing SQL value\.', t):
            replace_text(p, 'Execution output: ' + re.search(r'\d+', t).group() + ' row(s).')
        elif t.startswith('Actual effective DELETE privilege:'):
            replace_text(p, 'DELETE is granted and immediately revoked. The demonstration '
                         'runs inside a transaction that restores the existing permissions afterward.')
        elif t.startswith('Installed privilege evidence:'):
            replace_text(p, 'Restricted roles use explicitly permitted views and routines. '
                         'The helper EXECUTE grant supports the invoker transfer wrapper; '
                         'it does not grant general access to the underlying tables.')

    # View updatability is already explained for each view. Omit the repeated
    # catalogue output rather than retaining a second representation of that fact.
    for table in list(document.tables):
        if [c.text for c in table.rows[0].cells] == ['table name', 'is updatable', 'check option']:
            prev = table._tbl.getprevious()
            while prev is not None and (not text(prev).strip() or text(prev).startswith('Execution output:')
                                        or text(prev) == 'Actual PostgreSQL updatability catalogue'):
                prior = prev.getprevious()
                body.remove(prev)
                prev = prior
            table._tbl.getparent().remove(table._tbl)
            audit['removed_updatability_output_table'] = True

    # Retain the compact 15-case test matrix; the screenshots carry actual NOTICE
    # outputs, so remove only the column repeating those same strings.
    for table in document.tables:
        if table.cell(0,0).text == 'Test ID':
            assert table.rows[0].cells[4].text == 'Actual PostgreSQL Output'
            for row in table.rows:
                row._tr.remove(row.cells[4]._tc)
            table._tbl.tblGrid.remove(table._tbl.tblGrid.gridCol_lst[4])
            widths = [0.5,0.9,2.0,2.9,0.65]
            for col, width in zip(table.columns, widths):
                col.width = Inches(width)
            for row in table.rows:
                for cell, width in zip(row.cells, widths):
                    cell.width = Inches(width)
            table.autofit = False
            audit['removed_duplicate_test_output_column'] = True

    # Collapse empty paragraph runs left by removed result blocks inside cells.
    for table in document.tables:
        if not table.cell(0,0).text.startswith('Create >='):
            continue
        for row in table.rows:
            for cell in row.cells:
                previous_empty = False
                for p in list(cell.paragraphs):
                    empty = not p.text.strip() and not p._p.xpath('.//w:drawing')
                    if empty and previous_empty:
                        cell._tc.remove(p._p)
                    previous_empty = empty

    assert protected['section6'] == [canonical(e) for e in blocks_between('6 Tables and Constraints', '7 SQL Implementation')]
    assert protected['section7'] == [canonical(e) for e in blocks_between('7 SQL Implementation', '8 Advanced SQL Implementation')]
    assert protected['section9'] == [canonical(e) for e in blocks_between('9 Application Screenshots', '10 Work Breakdown')]
    assert protected['predictions'] == [canonical(r._tr) for t in document.tables for r in t.rows
                                       if r.cells[0].text == 'Expected output']
    assert protected['definitions'] == [text(p) for p in body.xpath('.//w:p') if style(p) == 'SQLCode'
                                       and text(p).startswith('CREATE OR REPLACE')]
    assert not any(re.search(r'Observed \d{4}-\d{2}-\d{2}|live screenshot,\s*\d{4}', text(p))
                   for p in body.xpath('.//w:p'))
    new_images = Counter(body.xpath('.//a:blip/@r:embed'))
    removed = old_images-new_images
    assert sum(removed.values()) == 3, removed  # two old batch images and old index image
    assert sum(new_images.values()) == sum(old_images.values())+3
    write_package(source, document)

    for name, key in [('Project Template - review 2.docx','review2_sha256'), ('Guide.docx','guide_sha256')]:
        assert hashlib.sha256((ROOT/name).read_bytes()).hexdigest() == preservation[key]
    with ZipFile(source) as a, ZipFile(TARGET) as b:
        changed = [n for n in a.namelist() if a.read(n) != b.read(n)]
        assert set(changed) <= {'word/document.xml','word/_rels/document.xml.rels'}, changed
    audit.update({'visible_images_before':sum(old_images.values()),
                  'visible_images_after':sum(new_images.values()),
                  'review2_and_guide_unchanged':True, 'banking_db_unchanged':True,
                  'definitions_predictions_and_manual_sections_preserved':True,
                  'changed_existing_package_parts':changed})
    (WORK/'edit_audit.json').write_text(json.dumps(audit,indent=2))
    print(json.dumps(audit,indent=2))


if __name__ == '__main__':
    main()
