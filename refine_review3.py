"""Refine the current, manually reviewed Review 3 without rebuilding it.

The current DOCX is authoritative. This patches document.xml only, preserving
all other package parts and every existing image. Review 2 is never written.
Run with --refresh-toc PDF after rendering to refresh the static contents page.
"""
from collections import Counter
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
import sys
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

ROOT=Path(__file__).resolve().parent
TARGET=ROOT/'Project Template - review 3.docx'
OTHER=ROOT/'Project Template - review 2.docx'
WORK=ROOT/'.report_work'/'review3-output-cleanup'
WORK.mkdir(parents=True,exist_ok=True)


def text(element):
    return ''.join(element.xpath('.//w:t/text()'))


def style(element):
    values=element.xpath('./w:pPr/w:pStyle/@w:val')
    return values[0] if values else ''


def replace_span(paragraph,start,end,replacement=''):
    """Replace text across split Word runs without resetting their formatting."""
    offset=0
    inserted=False
    for node in paragraph.xpath('.//w:t'):
        original=node.text or ''
        left=max(0,start-offset)
        right=min(len(original),end-offset)
        if right>left:
            node.text=original[:left]+(replacement if not inserted else '')+original[right:]
            inserted=True
        offset+=len(original)


def replace_text(paragraph,value):
    replace_span(paragraph,0,len(text(paragraph)),value)


def screenshot_follows(paragraph):
    node=paragraph.getnext()
    while node is not None:
        if node.tag==qn('w:p') and style(node).startswith('Heading'):
            return False
        if node.xpath('.//w:drawing'):
            return True
        node=node.getnext()
    return False


def drop_output_payload(paragraph,removed):
    node=paragraph.getnext()
    while node is not None:
        following=node.getnext()
        if node.xpath('.//w:drawing'):
            break
        is_table=node.tag==qn('w:tbl')
        is_result=node.tag==qn('w:p') and style(node)=='ResultCode'
        is_empty=node.tag==qn('w:p') and not text(node).strip()
        is_empty_result=node.tag==qn('w:p') and text(node).startswith('0 rows returned. The query executed successfully;')
        if not (is_table or is_result or is_empty or is_empty_result):
            break
        removed.append({'kind':'output table' if is_table else 'output paragraph','text':text(node)[:250]})
        node.getparent().remove(node)
        node=following


def protected_content(document):
    predictions=[]
    for table in document.tables:
        for row in table.rows:
            if row.cells[0].text.strip()=='Expected output':
                predictions.append(etree.tostring(row._tr,method='c14n'))
    return {
        'sql':[text(p) for p in document.element.body.xpath('.//w:p') if style(p)=='SQLCode'],
        'predictions':predictions,
        'images':document.element.body.xpath('.//a:blip/@r:embed'),
        'headings':[p.text for p in document.paragraphs if p.style.name.startswith('Heading')],
    }


def save_document_xml(document):
    temporary=TARGET.with_suffix('.refined.docx')
    xml=etree.tostring(document.element,xml_declaration=True,encoding='UTF-8',standalone=True)
    with ZipFile(TARGET) as source,ZipFile(temporary,'w') as output:
        for info in source.infolist():
            output.writestr(info,xml if info.filename=='word/document.xml' else source.read(info.filename))
    temporary.replace(TARGET)


def refresh_toc(document,pdf_path):
    from pypdf import PdfReader
    pages=[p.extract_text() for p in PdfReader(pdf_path).pages]
    normalize=lambda s:re.sub(r'\s+',' ',s).strip().casefold()
    headings={p.text for p in document.paragraphs if p.style.name.startswith('Heading')}
    mapping={}
    removed=[]
    for p in list(document.paragraphs):
        if p.style.name!='Report TOC':
            continue
        heading=p.text.rsplit('\t',1)[0]
        if heading not in headings:
            removed.append(heading)
            p._p.getparent().remove(p._p)
            continue
        matches=[i+1 for i,page in enumerate(pages) if any(normalize(line)==normalize(heading) for line in page.splitlines())]
        if not matches:
            raise RuntimeError('Cannot locate heading: '+heading)
        mapping[heading]=matches[-1]
        replace_text(p._p,f'{heading}\t{matches[-1]}')
    (WORK/'page_map.json').write_text(json.dumps(mapping,indent=2),encoding='utf-8')
    return {'pages':len(pages),'removed_stale_contents_entries':removed,'page_map':mapping}


def main():
    other_hash=hashlib.sha256(OTHER.read_bytes()).hexdigest()
    document=Document(TARGET)
    before=protected_content(document)
    removed=[]
    counts={'captions_cleaned':0,'output_previews_simplified':0,'proof_tables_removed':0,'indexed_plan_text_removed':0}
    if '--refresh-toc' in sys.argv:
        details=refresh_toc(document,Path(sys.argv[sys.argv.index('--refresh-toc')+1]))
    else:
        for p in document.element.body.xpath('.//w:p'):
            suffix=re.search(r'\s*[—–-]\s*live screenshot,\s*\d{4}-\d{2}-\d{2}\.?\s*$',text(p))
            if suffix:
                replace_span(p,suffix.start(),suffix.end())
                counts['captions_cleaned']+=1
        for p in list(document.element.body.xpath('.//w:p')):
            match=re.fullmatch(r'Actual output: (\d+) row\(s\); first \d+ shown\. NULL denotes a missing SQL value\.',text(p))
            if not match or not screenshot_follows(p):
                continue
            drop_output_payload(p,removed)
            # Counts refer to the recorded execution. We do not claim that the
            # old typed three-row sample equals the rows visible in a viewport.
            replace_text(p,f'Execution output: {match[1]} row(s).')
            counts['output_previews_simplified']+=1
        for table in list(document.tables):
            header=[c.text.strip() for c in table.rows[0].cells]
            if header not in [['Observation','Actual value'],['Session','Actual observation']]:
                continue
            if screenshot_follows(table._tbl):
                removed.append({'kind':'duplicate proof table','text':text(table._tbl)[:250]})
                table._tbl.getparent().remove(table._tbl)
                counts['proof_tables_removed']+=1
        for p in document.paragraphs:
            if p.text=='Actual plan with idx_txn_account_date available' and screenshot_follows(p._p):
                node=p._p.getnext()
                if node is not None and style(node)=='ResultCode':
                    removed.append({'kind':'indexed plan duplicate','text':text(node)[:250]})
                    node.getparent().remove(node)
                    counts['indexed_plan_text_removed']+=1
        # Remove stale TOC entries left behind by the user's own final deletions.
        headings={p.text for p in document.paragraphs if p.style.name.startswith('Heading')}
        stale=[]
        for p in list(document.paragraphs):
            if p.style.name=='Report TOC' and p.text.rsplit('\t',1)[0] not in headings:
                stale.append(p.text)
                p._p.getparent().remove(p._p)
        details={'changes':counts,'removed_payloads':removed,'removed_stale_contents_entries':stale}
    assert protected_content(document)==before,'A protected heading, SQL block, prediction or image changed'
    save_document_xml(document)
    assert protected_content(Document(TARGET))==before
    assert hashlib.sha256(OTHER.read_bytes()).hexdigest()==other_hash,'Review 2 changed'
    mode='toc' if '--refresh-toc' in sys.argv else 'cleanup'
    (WORK/(mode+'-audit.json')).write_text(json.dumps(details,indent=2),encoding='utf-8')
    print(json.dumps(details if mode=='toc' else {'changes':counts,'removed_payload_count':len(removed),'review2_unchanged':True,'protected_content_preserved':True},indent=2))


if __name__=='__main__':
    main()
